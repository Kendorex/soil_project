from __future__ import annotations

import csv
import json
import math
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .apsim_editor import _named
from .apsim_runner import find_db_for_apsimx, run_apsim_files
from .cultivar_calibration import (
    CalibrationError,
    create_calibration_copy,
    load_cultivar_observations,
)

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


@dataclass(frozen=True)
class WaterObservation:
    year: int
    field: int
    measurement_date: date
    gravimetric_percent_mean: float
    bulk_density_mean_g_cm3: float
    volumetric_fraction_target: float
    point_count: int


def _parse_day_month(value: Any, year: int) -> date:
    if isinstance(value, datetime):
        return value.date().replace(year=year)
    if isinstance(value, date):
        return value.replace(year=year)
    text = str(value).strip().replace(",", ".")
    parts = text.split(".")
    if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
        return date(year, int(parts[1]), int(parts[0]))
    raise CalibrationError(f"Cannot parse soil-water measurement date {value!r} for {year}")


def gravimetric_to_volumetric_fraction(moisture_percent: float, bulk_density_g_cm3: float) -> float:
    """Convert gravimetric water percent to m3/m3 using rho_water=1 g/cm3."""
    return float(moisture_percent) / 100.0 * float(bulk_density_g_cm3)


def load_water_observations(config: dict[str, Any]) -> list[WaterObservation]:
    if load_workbook is None:
        raise CalibrationError("openpyxl is required to read soil-water observations")
    root = config["_root"]
    calibration = config["calibration"]
    water_config = config["initial_water_calibration"]
    workbook_path = root / calibration["observations_workbook"]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        moisture_sheet = workbook[str(water_config["measurement_sheet"])]
        density_sheet = workbook[str(water_config["density_sheet"])]
        moisture_rows = list(moisture_sheet.iter_rows(values_only=True))[1:]
        density_rows = list(density_sheet.iter_rows(values_only=True))[1:]

        density_by_point: dict[tuple[int, int, int], float] = {}
        for row in density_rows:
            try:
                density_by_point[(int(row[0]), int(row[1]), int(row[2]))] = float(row[3])
            except (TypeError, ValueError):
                continue

        field_by_year = {item.year: item.field for item in load_cultivar_observations(config)}
        observations = []
        for raw_year, raw_date in water_config["measurement_dates"].items():
            year = int(raw_year)
            field = field_by_year[year]
            target_date = _parse_day_month(raw_date, year)
            paired = []
            for row in moisture_rows:
                try:
                    row_year, row_field, point = int(row[0]), int(row[1]), int(row[2])
                    row_date = _parse_day_month(row[3], row_year)
                    moisture = float(row[4])
                except (TypeError, ValueError, CalibrationError):
                    continue
                if (row_year, row_field, row_date) != (year, field, target_date):
                    continue
                density = density_by_point.get((year, field, point))
                if density is not None:
                    paired.append((moisture, density))
            if not paired:
                raise CalibrationError(
                    f"No paired moisture/density measurements for year={year}, field={field}, date={target_date}"
                )
            observations.append(
                WaterObservation(
                    year=year,
                    field=field,
                    measurement_date=target_date,
                    gravimetric_percent_mean=sum(item[0] for item in paired) / len(paired),
                    bulk_density_mean_g_cm3=sum(item[1] for item in paired) / len(paired),
                    volumetric_fraction_target=sum(
                        gravimetric_to_volumetric_fraction(*item) for item in paired
                    ) / len(paired),
                    point_count=len(paired),
                )
            )
        return sorted(observations, key=lambda item: item.year)
    finally:
        workbook.close()


def water_between_airdry_and_dul(airdry: float, dul: float, relative_fraction: float) -> float:
    if not 0.0 <= relative_fraction <= 1.0:
        raise CalibrationError("initial AirDry-to-DUL fraction must be between 0 and 1")
    return float(airdry) + float(relative_fraction) * (float(dul) - float(airdry))


def patch_initial_water(apsimx_path: Path, relative_fraction: float, layers_to_adjust: int) -> list[float]:
    data = json.loads(apsimx_path.read_text(encoding="utf-8"))
    water = _named(data, "Water")
    physical = _named(data, "Physical")
    values = list(water.get("InitialValues") or [])
    airdry = list(physical.get("AirDry") or [])
    dul = list(physical.get("DUL") or [])
    if not values or len(values) != len(airdry) or len(values) != len(dul):
        raise CalibrationError("Soil Water.InitialValues, Physical.AirDry and Physical.DUL must have equal layers")
    if layers_to_adjust < 1 or layers_to_adjust > len(values):
        raise CalibrationError(f"layers_to_adjust must be in 1..{len(values)}")
    for index in range(layers_to_adjust):
        values[index] = water_between_airdry_and_dul(airdry[index], dul[index], relative_fraction)
    water["InitialValues"] = values
    apsimx_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return values


def _fixed_cultivar_parameters(config: dict[str, Any]) -> dict[str, float]:
    parameters = {}
    for name, item in config["calibration"].get("parameter_grid", {}).items():
        values = list(item.get("values", []))
        if len(values) != 1:
            raise CalibrationError(
                f"Initial-water calibration requires one fixed cultivar value for {name}; found {values}"
            )
        parameters[name] = float(values[0])
    return parameters


def build_initial_water_plan(
    optimizer_config: dict[str, Any], config: dict[str, Any], stage_files: bool = True
) -> list[dict[str, Any]]:
    crop_observations = {item.year: item for item in load_cultivar_observations(config)}
    water_observations = load_water_observations(config)
    water_config = config["initial_water_calibration"]
    fractions = [float(value) for value in water_config["airdry_to_dul_fractions"]]
    root = config["_root"]
    work_dir = root / water_config["work_dir"]
    parameters = _fixed_cultivar_parameters(config)
    plan = []
    for observation in water_observations:
        crop = crop_observations[observation.year]
        for fraction in fractions:
            fraction_label = f"water_{fraction:.3f}".replace(".", "p")
            destination = work_dir / str(observation.year) / fraction_label / f"zlata_{observation.year}.apsimx"
            initial_values = []
            if stage_files:
                create_calibration_copy(optimizer_config, config, crop, parameters, destination)
                initial_values = patch_initial_water(
                    destination, fraction, int(water_config["layers_to_adjust"])
                )
            plan.append(
                {
                    "year": observation.year,
                    "field": observation.field,
                    "candidate_id": f"{observation.year}_{fraction_label}",
                    "initial_airdry_to_dul_fraction": fraction,
                    "layers_adjusted": int(water_config["layers_to_adjust"]),
                    "measurement_date": observation.measurement_date.isoformat(),
                    "measured_gravimetric_percent_mean": observation.gravimetric_percent_mean,
                    "bulk_density_mean_g_cm3": observation.bulk_density_mean_g_cm3,
                    "target_volumetric_fraction": observation.volumetric_fraction_target,
                    "measurement_point_count": observation.point_count,
                    "moisture_units_assumption": water_config["moisture_units_assumption"],
                    "working_apsimx": destination.relative_to(root).as_posix(),
                    "initial_water_values": " | ".join(f"{value:.6f}" for value in initial_values),
                }
            )
    return plan


def _parse_sqlite_date(value: Any) -> date | None:
    if value is None:
        return None
    text = str(value).strip()
    for candidate in (text, text.split(" ", 1)[0]):
        try:
            return datetime.fromisoformat(candidate).date()
        except ValueError:
            continue
    return None


def read_simulated_topsoil_water(database: Path, measurement_date: date) -> dict[str, float]:
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    try:
        columns = [str(row[1]) for row in connection.execute('PRAGMA table_info("Report")')]
        date_column = next((name for name in ("Clock.Today", "Date", "Today") if name in columns), None)
        if date_column is None or "SW(1)" not in columns or "SW(2)" not in columns:
            raise CalibrationError(f"Report needs date, SW(1), SW(2); available={columns}")
        rows = connection.execute(
            f'SELECT "{date_column}", "SW(1)", "SW(2)" FROM "Report"'
        ).fetchall()
        selected = next((row for row in rows if _parse_sqlite_date(row[0]) == measurement_date), None)
        if selected is None:
            raise CalibrationError(f"Report has no row for soil-water measurement date {measurement_date}")
        sw1, sw2 = float(selected[1]), float(selected[2])
        # APSIM layers are 0-15 and 15-30 cm. The observed 0-20 cm value is
        # represented by 15 cm of layer 1 plus 5 cm of layer 2.
        weighted = (15.0 * sw1 + 5.0 * sw2) / 20.0
        return {"simulated_sw_0_15": sw1, "simulated_sw_15_30": sw2, "simulated_sw_0_20": weighted}
    finally:
        connection.close()


def run_initial_water_plan(
    plan: list[dict[str, Any]], optimizer_config: dict[str, Any], config: dict[str, Any], logger
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    root = config["_root"]
    results = []
    for planned in plan:
        result = dict(planned)
        try:
            apsimx = root / planned["working_apsimx"]
            run_apsim_files([apsimx], optimizer_config, logger)
            database = find_db_for_apsimx(apsimx)
            if database is None:
                raise CalibrationError(f"APSIMX database not found for {apsimx}")
            simulated = read_simulated_topsoil_water(
                database, date.fromisoformat(planned["measurement_date"])
            )
            error = simulated["simulated_sw_0_20"] - float(planned["target_volumetric_fraction"])
            result.update(simulated)
            result.update(
                {
                    "error_volumetric_fraction": error,
                    "absolute_error_volumetric_fraction": abs(error),
                    "run_status": "completed",
                    "error_message": "",
                }
            )
        except Exception as exc:
            result.update(
                {
                    "simulated_sw_0_15": None,
                    "simulated_sw_15_30": None,
                    "simulated_sw_0_20": None,
                    "error_volumetric_fraction": None,
                    "absolute_error_volumetric_fraction": None,
                    "run_status": "failed",
                    "error_message": str(exc),
                }
            )
            logger.exception(
                "Initial-water calibration failed: year=%s fraction=%s",
                planned["year"], planned["initial_airdry_to_dul_fraction"],
            )
        results.append(result)

    best = []
    for year in sorted({int(row["year"]) for row in results}):
        valid = [
            row for row in results
            if int(row["year"]) == year and row["absolute_error_volumetric_fraction"] is not None
        ]
        if valid:
            best.append(min(valid, key=lambda row: (
                float(row["absolute_error_volumetric_fraction"]),
                float(row["initial_airdry_to_dul_fraction"]),
            )))
    return results, best


def write_rows(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
