#!/usr/bin/env python3
"""Export and summarize supplied dry-biomass observations for oat fields."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import sqlite3
import tarfile
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OBSERVATIONS = ROOT / "data" / "observations" / "multifield_2019_2023.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or ["empty"])
        writer.writeheader()
        writer.writerows(rows)


def create_archive(path: Path, output: Path) -> None:
    files = sorted(file for file in output.iterdir() if file.is_file())
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [{
            "path": file.name,
            "size": file.stat().st_size,
            "sha256": sha256(file),
        } for file in files],
    }
    with tarfile.open(path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.name, recursive=False)


def load_source_biomass_rows(observations: dict) -> tuple[list[dict], dict]:
    """Read the source biomass table after verifying workbook provenance."""
    workbook_path = Path(observations["source"]["workbook"])
    workbook_hash = sha256(workbook_path)
    if workbook_hash != observations["source"].get("workbook_sha256"):
        raise ValueError("source workbook hash does not match observations metadata")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Сухая_биомасса"]
    values = sheet.iter_rows(values_only=True)
    header = [str(value) if value not in (None, "") else f"column_{index}"
              for index, value in enumerate(next(values), start=1)]
    rows = [
        {"source_row": number, **dict(zip(header, row))}
        for number, row in enumerate(values, start=2)
        if any(value not in (None, "") for value in row)
    ]
    sheet_name = sheet.title
    workbook.close()
    return rows, {
        "workbook": str(workbook_path),
        "workbook_sha256": workbook_hash,
        "sheet": sheet_name,
        "columns": header,
    }


def model_phase_comparison(database: Path, observations: dict,
                           source_rows: list[dict]) -> list[dict]:
    phase_map = {
        "кущение": "observed_tillering_date",
        "выметывание": "observed_heading_date",
    }
    records = [
        (int(year), row)
        for year, year_rows in observations["years"].items()
        if 2019 <= int(year) <= 2023
        for row in year_rows
        if "Oats" in row.get("yield_components", [])
    ]
    result = []
    uri = database.resolve().as_uri() + "?mode=ro"
    with sqlite3.connect(uri, uri=True) as connection:
        connection.row_factory = sqlite3.Row
        available_tables = {
            row[0] for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )
        }
        simulation = connection.execute(
            'SELECT ID, Name FROM "_Simulations" ORDER BY ID LIMIT 1'
        ).fetchone()
        for year, record in records:
            field = record["field"]
            table = f"Report{field.removeprefix('Field')}"
            if table not in available_tables:
                continue
            for source_phase, date_key in phase_map.items():
                observed = [
                    float(row["Сухая биомасса, г/м2"])
                    for row in source_rows
                    if row.get("Год") == year
                    and row.get("Поле") == int(field.removeprefix("Field"))
                    and str(row.get("Фенофаза (сухая биомасса)")).casefold() == source_phase
                    and isinstance(row.get("Сухая биомасса, г/м2"), (int, float))
                ]
                date = record.get(date_key)
                if not observed or not date:
                    continue
                model = connection.execute(
                    f'''SELECT OatsAGWt, CloverAGWt, PlantainAGWt, MixAGWt,
                               OatsTotalWt, CloverTotalWt, PlantainTotalWt
                        FROM "{table}"
                        WHERE SimulationID=? AND date("Clock.Today")=date(?)''',
                    (simulation["ID"], date),
                ).fetchone()
                if model is None:
                    raise ValueError(
                        f"missing model comparison data for {field} {year} {source_phase}"
                    )
                value = lambda name: float(model[name] or 0.0)
                oats = value("OatsAGWt")
                clover = value("CloverAGWt")
                plantain = value("PlantainAGWt")
                total_stand = oats + clover + plantain
                mean = fmean(observed)
                result.append({
                    "field": field,
                    "year": year,
                    "cultivar": record.get("cultivar"),
                    "source_phase": source_phase,
                    "comparison_date": date,
                    "observed_n": len(observed),
                    "observed_mean_g_m2": mean,
                    "observed_sd_g_m2": pstdev(observed),
                    "observed_min_g_m2": min(observed),
                    "observed_max_g_m2": max(observed),
                    "model_oats_agwt_g_m2": oats,
                    "model_clover_agwt_g_m2": clover,
                    "model_plantain_agwt_g_m2": plantain,
                    "model_total_stand_agwt_g_m2": total_stand,
                    "model_mix_agwt_g_m2": value("MixAGWt"),
                    "model_oats_total_wt_g_m2": value("OatsTotalWt"),
                    "model_clover_total_wt_g_m2": value("CloverTotalWt"),
                    "model_plantain_total_wt_g_m2": value("PlantainTotalWt"),
                    "oats_only_error_g_m2": oats - mean,
                    "total_stand_error_g_m2": total_stand - mean,
                    "simulation": simulation["Name"],
                })
    return result


def comparison_metrics(rows: list[dict]) -> dict:
    scopes = {
        "all_oats_2019_2023": rows,
        "yakov_development_2020_2022": [
            row for row in rows if 2020 <= row["year"] <= 2022
            and str(row.get("cultivar", "")).casefold() == "яков"
        ],
        "yakov_validation_2023": [
            row for row in rows if row["year"] == 2023
            and str(row.get("cultivar", "")).casefold() == "яков"
        ],
    }
    result = {}
    for scope, selected in scopes.items():
        errors = [row["oats_only_error_g_m2"] for row in selected]
        result[scope] = {
            "n": len(selected),
            "mae_g_m2": fmean(abs(value) for value in errors) if errors else None,
            "rmse_g_m2": (fmean(value * value for value in errors) ** 0.5) if errors else None,
            "bias_g_m2": fmean(errors) if errors else None,
        }
    return result


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--observations", type=Path, default=DEFAULT_OBSERVATIONS)
    parser.add_argument("--database", type=Path)
    args = parser.parse_args(argv)
    observations = json.loads(args.observations.read_text(encoding="utf-8"))
    try:
        rows, source = load_source_biomass_rows(observations)
    except ValueError as error:
        parser.error(str(error))

    years = sorted({str(row.get("Год")) for row in rows if row.get("Год") is not None})
    fields = sorted({str(row.get("Поле")) for row in rows if row.get("Поле") is not None})
    oat_2023 = [row for row in rows
                if row.get("Год") == 2023 and row.get("Поле") in (2, 7)]

    comparison = (
        model_phase_comparison(args.database, observations, rows)
        if args.database else []
    )
    metrics = comparison_metrics(comparison) if comparison else {}

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output = ROOT / "outputs" / "multifield" / "OATS_BIOMASS_SOURCE_AUDIT" / stamp
    output.mkdir(parents=True, exist_ok=False)
    write_csv(output / "source_dry_biomass_rows.csv", rows)
    write_csv(output / "oats_2023_field2_field7_rows.csv", oat_2023)
    if comparison:
        write_csv(output / "model_phase_comparison.csv", comparison)
    summary = {
        "experiment": "source_dry_biomass_observation_audit",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": "input_evidence_audit",
        "base_model_changed": False,
        **source,
        "row_count": len(rows),
        "years": years,
        "fields": fields,
        "oats_2023_field2_field7_row_count": len(oat_2023),
        "database": str(args.database.resolve()) if args.database else None,
        "database_sha256": sha256(args.database) if args.database else None,
        "model_phase_comparison": comparison,
        "comparison_metrics": metrics,
    }
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    archive = ROOT / f"oats_biomass_source_audit_{stamp}.tar.gz"
    create_archive(archive, output)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"OUTPUT: {output}")
    print(f"ARCHIVE: {archive}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
