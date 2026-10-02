from __future__ import annotations

import csv
import json
import math
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.config import load_config, load_inputs
from src.cultivar_calibration import load_calibration_config, load_cultivar_observations


OUTPUT = ROOT / "docs" / "Отчет_по_калибровке_сортов_Zlata_Irgina.docx"
CALIBRATION_RESULTS = ROOT / "examples" / "final_calibration_results"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "243447"
MUTED = "667085"
LIGHT = "F2F4F7"
CALLOUT = "EAF2F8"
CAUTION = "FFF4CE"
WHITE = "FFFFFF"


def _font(run, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
    else:
        tc_pr.remove(shd)
    shd.set(qn("w:fill"), fill)
    # CT_TcPr requires shading before margins and vertical alignment.
    insert_at = len(tc_pr)
    for index, child in enumerate(tc_pr):
        if child.tag in {qn("w:tcMar"), qn("w:textDirection"), qn("w:tcFitText"), qn("w:vAlign"), qn("w:hideMark")}:
            insert_at = index
            break
    tc_pr.insert(insert_at, shd)


def _cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    # Word's transitional OOXML expects left/right in tcMar.
    for tag, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
    else:
        tbl_pr.remove(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    # CT_TblPr requires tblInd after justification and before tblLayout/tblLook.
    insert_at = len(tbl_pr)
    for index, child in enumerate(tbl_pr):
        if child.tag in {qn("w:tblBorders"), qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"), qn("w:tblLook")}:
            insert_at = index
            break
    tbl_pr.insert(insert_at, tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths_dxa[index])
            _cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _format_table(table, header=True, font_size=9.3):
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if header and row_index == 0:
                _shade(cell, LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    _font(run, font_size, bold=(header and row_index == 0), color=INK)


def _add_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int], font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    _set_table_geometry(table, widths_dxa)
    _format_table(table, font_size=font_size)
    return table


def _keep_table_together(table):
    """Keep a compact table on one page when Word can fit it."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def _add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def _add_body(doc, text: str, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        _font(paragraph.add_run(bold_lead), bold=True, color=INK)
        _font(paragraph.add_run(text[len(bold_lead):]), color=INK)
    else:
        _font(paragraph.add_run(text), color=INK)
    return paragraph


def _add_callout(doc, label: str, text: str, fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    _set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    _shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    _font(p.add_run(label + " "), bold=True, color=DARK_BLUE)
    _font(p.add_run(text), color=INK)
    return table


def _extract_cultivar_parameters(template: Path) -> dict[str, dict[str, float]]:
    data = json.loads(template.read_text(encoding="utf-8"))
    wanted = {
        "MinLN": "[Phenology].CAMP.FLNparams.MinLN",
        "PpLN": "[Phenology].CAMP.FLNparams.PpLN",
        "BasePhyllochron": "[Phenology].Phyllochron.BasePhyllochron.FixedValue",
        "GrainFillingTarget": "[Phenology].GrainFilling.Target.FixedValue",
        "GrainRipeningTarget": "[Phenology].GrainRipening.Target.FixedValue",
    }

    def walk(value):
        if isinstance(value, dict):
            yield value
            for child in value.values():
                yield from walk(child)
        elif isinstance(value, list):
            for child in value:
                yield from walk(child)

    result = {}
    for node in walk(data):
        name = node.get("Name")
        if name not in {"Zlata", "Irgina"}:
            continue
        commands = node.get("Command", node.get("Commands", []))
        values = {}
        for label, expression in wanted.items():
            for command in commands:
                if str(command).strip().startswith(expression):
                    match = re.search(r"=\s*([-+]?\d+(?:\.\d+)?)", str(command))
                    if match:
                        values[label] = float(match.group(1))
                        break
        if values:
            result[name] = values
    return result


def _load_observations(config_name: str):
    cfg = load_calibration_config(ROOT / "configs" / config_name, ROOT)
    return load_cultivar_observations(cfg)


def _load_final_results(cultivar_slug: str):
    folder = CALIBRATION_RESULTS / f"cultivar_calibration_{cultivar_slug}_final_evaluation"
    runs_path = folder / "calibration_runs.csv"
    summary_path = folder / "calibration_summary.csv"
    if not runs_path.is_file() or not summary_path.is_file():
        raise RuntimeError(f"Final calibration outputs are missing: {folder}")
    with runs_path.open(encoding="utf-8-sig", newline="") as stream:
        runs = list(csv.DictReader(stream))
    with summary_path.open(encoding="utf-8-sig", newline="") as stream:
        summaries = list(csv.DictReader(stream))
    if len(summaries) != 1:
        raise RuntimeError(f"Expected one final candidate summary in {summary_path}, found {len(summaries)}")
    return runs, summaries[0]


def _iso_doy(value: str) -> int:
    return date.fromisoformat(value).timetuple().tm_yday


def _yield_t_ha_from_run(row: dict[str, str]) -> float:
    raw = float(row["simulated_yield_raw"] or 0)
    return raw / 1000.0


def _yield_metrics(rows: list[dict[str, str]]) -> dict[str, float]:
    observed = [float(row["observed_yield_t_ha"]) for row in rows]
    simulated = [_yield_t_ha_from_run(row) for row in rows]
    errors = [sim - obs for obs, sim in zip(observed, simulated)]
    return {
        "observed_mean": sum(observed) / len(observed),
        "simulated_mean": sum(simulated) / len(simulated),
        "mae": sum(abs(value) for value in errors) / len(errors),
        "rmse": math.sqrt(sum(value * value for value in errors) / len(errors)),
        "bias": sum(errors) / len(errors),
    }


def _image_font(size: int, bold: bool = False):
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [windows / ("arialbd.ttf" if bold else "arial.ttf"), windows / "calibri.ttf"]
    for path in candidates:
        if path.is_file():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def _text(draw, xy, value, size=24, fill="#243447", bold=False, anchor=None):
    draw.text(xy, str(value), font=_image_font(size, bold), fill=fill, anchor=anchor)


def _scale(value: float, minimum: float, maximum: float, low: float, high: float) -> float:
    if maximum == minimum:
        return (low + high) / 2
    return low + (value - minimum) / (maximum - minimum) * (high - low)


def _draw_line_panel(draw, box, title, years, series, y_label):
    left, top, right, bottom = box
    plot_left, plot_top, plot_right, plot_bottom = left + 105, top + 65, right - 25, bottom - 75
    values = [value for _, data, _ in series for value in data if not math.isnan(value)]
    y_min = math.floor((min(values) - 8) / 10) * 10
    y_max = math.ceil((max(values) + 8) / 10) * 10
    _text(draw, ((left + right) / 2, top + 10), title, 28, bold=True, anchor="ma")
    _text(draw, (left + 4, plot_top - 28), y_label, 18, fill="#667085")
    for step in range(6):
        value = y_min + step * (y_max - y_min) / 5
        y = _scale(value, y_min, y_max, plot_bottom, plot_top)
        draw.line((plot_left, y, plot_right, y), fill="#D0D5DD", width=1)
        _text(draw, (plot_left - 12, y), f"{value:.0f}", 17, fill="#667085", anchor="rm")
    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#98A2B3", width=2)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#98A2B3", width=2)
    x_positions = []
    for index, year in enumerate(years):
        x = _scale(index, 0, max(1, len(years) - 1), plot_left, plot_right)
        x_positions.append(x)
        _text(draw, (x, plot_bottom + 18), year, 16, fill="#475467", anchor="ma")
    for label, data, color in series:
        points = []
        for x, value in zip(x_positions, data):
            if math.isnan(value):
                continue
            y = _scale(value, y_min, y_max, plot_bottom, plot_top)
            points.append((x, y))
        if len(points) > 1:
            draw.line(points, fill=color, width=4, joint="curve")
        for x, y in points:
            draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill=color, outline="white", width=2)
    legend_x = plot_left
    legend_y = bottom - 30
    for label, _, color in series:
        draw.line((legend_x, legend_y, legend_x + 34, legend_y), fill=color, width=5)
        _text(draw, (legend_x + 44, legend_y), label, 17, anchor="lm")
        legend_x += 280


def _plot_heading_comparison(zlata_runs, irgina_runs, destination: Path):
    image = Image.new("RGB", (1800, 1300), "white")
    draw = ImageDraw.Draw(image)
    for box, rows, cultivar in (
        ((35, 25, 1765, 625), zlata_runs, "Zlata"),
        ((35, 665, 1765, 1265), irgina_runs, "Irgina"),
    ):
        years = [int(row["year"]) for row in rows]
        observed = [_iso_doy(row["observed_heading_date"]) for row in rows]
        simulated = [_iso_doy(row["simulated_heading_date"]) for row in rows]
        _draw_line_panel(
            draw,
            box,
            f"{cultivar}: начало колошения",
            years,
            [("Наблюдение", observed, "#2E74B5"), ("APSIMX", simulated, "#C69214")],
            "День года",
        )
    image.save(destination, quality=95)


def _plot_maturity_comparison(zlata_runs, irgina_runs, destination: Path):
    image = Image.new("RGB", (1800, 1300), "white")
    draw = ImageDraw.Draw(image)
    for box, rows, cultivar in (
        ((35, 25, 1765, 625), zlata_runs, "Zlata"),
        ((35, 665, 1765, 1265), irgina_runs, "Irgina"),
    ):
        years = [int(row["year"]) for row in rows]
        harvest = [_iso_doy(row["observed_harvest_date"]) for row in rows]
        maturity = [_iso_doy(row["simulated_maturity_date"]) for row in rows]
        _draw_line_panel(
            draw,
            box,
            f"{cultivar}: зрелость и уборка",
            years,
            [("Фактическая уборка", harvest, "#667085"), ("Зрелость APSIMX", maturity, "#14866D")],
            "День года",
        )
    image.save(destination, quality=95)


def _plot_yield_comparison(zlata_runs, irgina_runs, destination: Path):
    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    all_yields = [
        value
        for rows in (zlata_runs, irgina_runs)
        for row in rows
        for value in (float(row["observed_yield_t_ha"]), _yield_t_ha_from_run(row))
    ]
    y_max = math.ceil((max(all_yields) + 0.4) * 2) / 2
    for box, rows, cultivar in (
        ((25, 20, 890, 735), zlata_runs, "Zlata"),
        ((910, 20, 1775, 735), irgina_runs, "Irgina"),
    ):
        left, top, right, bottom = box
        pl, pt, pr, pb = left + 90, top + 70, right - 20, bottom - 90
        years = [int(row["year"]) for row in rows]
        observed = [float(row["observed_yield_t_ha"]) for row in rows]
        simulated = [_yield_t_ha_from_run(row) for row in rows]
        _text(draw, ((left + right) / 2, top + 15), cultivar, 30, bold=True, anchor="ma")
        _text(draw, (left + 4, pt - 30), "Урожайность, т/га", 17, fill="#667085")
        for step in range(5):
            value = step * y_max / 4
            y = _scale(value, 0, y_max, pb, pt)
            draw.line((pl, y, pr, y), fill="#D0D5DD", width=1)
            _text(draw, (pl - 10, y), f"{value:.1f}", 16, fill="#667085", anchor="rm")
        draw.line((pl, pt, pl, pb), fill="#98A2B3", width=2)
        draw.line((pl, pb, pr, pb), fill="#98A2B3", width=2)
        slot = (pr - pl) / len(years)
        for index, (year, obs, sim) in enumerate(zip(years, observed, simulated)):
            center = pl + (index + 0.5) * slot
            bar_width = slot * 0.28
            for x0, value, color in ((center - bar_width, obs, "#2E74B5"), (center, sim, "#C69214")):
                x1 = x0 + bar_width
                y = _scale(value, 0, y_max, pb, pt)
                draw.rectangle((x0, y, x1, pb), fill=color)
            _text(draw, (center, pb + 17), year, 14, fill="#475467", anchor="ma")
        legend_x = pl
        for label, color in (("Наблюдение", "#2E74B5"), ("APSIMX", "#C69214")):
            draw.rectangle((legend_x, bottom - 36, legend_x + 24, bottom - 12), fill=color)
            _text(draw, (legend_x + 34, bottom - 24), label, 15, anchor="lm")
            legend_x += 210
    image.save(destination, quality=95)


def _plot_parameters(parameters: dict[str, dict[str, float]], destination: Path):
    image = Image.new("RGB", (1800, 780), "white")
    draw = ImageDraw.Draw(image)
    names_left = ["MinLN", "PpLN"]
    names_right = ["BasePhyllochron", "GrainFillingTarget", "GrainRipeningTarget"]
    colors = {"Zlata": "#2E74B5", "Irgina": "#C69214"}
    for box, names, title in (
        ((25, 20, 750, 755), names_left, "Коэффициенты фенологии"),
        ((780, 20, 1775, 755), names_right, "Тепловые параметры, °C·сут"),
    ):
        left, top, right, bottom = box
        pl, pt, pr, pb = left + 75, top + 75, right - 25, bottom - 105
        maximum = max(parameters[c][n] for c in ("Zlata", "Irgina") for n in names) * 1.18
        _text(draw, ((left + right) / 2, top + 15), title, 25, bold=True, anchor="ma")
        for step in range(5):
            value = step * maximum / 4
            y = _scale(value, 0, maximum, pb, pt)
            draw.line((pl, y, pr, y), fill="#D0D5DD", width=1)
            _text(draw, (pl - 10, y), f"{value:.0f}" if maximum > 20 else f"{value:.1f}", 15, fill="#667085", anchor="rm")
        draw.line((pl, pt, pl, pb), fill="#98A2B3", width=2)
        draw.line((pl, pb, pr, pb), fill="#98A2B3", width=2)
        slot = (pr - pl) / len(names)
        for index, name in enumerate(names):
            center = pl + (index + 0.5) * slot
            bar_width = slot * 0.28
            for offset, cultivar in ((-bar_width, "Zlata"), (0, "Irgina")):
                value = parameters[cultivar][name]
                x0 = center + offset
                x1 = x0 + bar_width
                y = _scale(value, 0, maximum, pb, pt)
                draw.rectangle((x0, y, x1, pb), fill=colors[cultivar])
                _text(draw, ((x0 + x1) / 2, y - 8), f"{value:g}", 16, bold=True, anchor="mb")
            label = name.replace("Target", "\nTarget") if "Target" in name else name
            for line_index, line in enumerate(label.split("\n")):
                _text(draw, (center, pb + 20 + line_index * 20), line, 14, fill="#475467", anchor="ma")
        legend_x = pl
        for cultivar in ("Zlata", "Irgina"):
            draw.rectangle((legend_x, bottom - 36, legend_x + 24, bottom - 12), fill=colors[cultivar])
            _text(draw, (legend_x + 34, bottom - 24), cultivar, 16, anchor="lm")
            legend_x += 150
    image.save(destination, quality=95)


def _configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    _font(hp.add_run("APSIMX · калибровка сортов"), 8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(fp.add_run("ВНИИМЗ · предварительный технический отчёт · "), 8, color=MUTED)
    run = fp.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def _add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    _font(p.add_run(text), 8.5, color=MUTED)


def build_report():
    zlata = _load_observations("cultivar_calibration_zlata_final_evaluation.yaml")
    irgina = _load_observations("cultivar_calibration_irgina_final_evaluation.yaml")
    zlata_runs, zlata_summary = _load_final_results("zlata")
    irgina_runs, irgina_summary = _load_final_results("irgina")
    zlata_yield_metrics = _yield_metrics(zlata_runs)
    irgina_yield_metrics = _yield_metrics(irgina_runs)
    parameters = _extract_cultivar_parameters(ROOT / "simulation" / "standart.apsimx")
    optimizer_config = load_config(ROOT / "configs" / "optimizer_config.yaml")
    optimizer_inputs = load_inputs(optimizer_config)
    economic_constants = optimizer_inputs["constants"]
    fertilizer_prices = optimizer_inputs["prices"]
    seed_config = optimizer_config.get("seed_economics", {})
    if set(parameters) != {"Zlata", "Irgina"}:
        raise RuntimeError(f"Cannot read both cultivar parameter sets: {parameters}")

    with tempfile.TemporaryDirectory(prefix="apsim_calibration_report_") as tmp:
        tmp = Path(tmp)
        phenology_chart = tmp / "heading_comparison.png"
        maturity_chart = tmp / "maturity_comparison.png"
        yield_chart = tmp / "yield_comparison.png"
        parameter_chart = tmp / "parameters.png"
        _plot_heading_comparison(zlata_runs, irgina_runs, phenology_chart)
        _plot_maturity_comparison(zlata_runs, irgina_runs, maturity_chart)
        _plot_yield_comparison(zlata_runs, irgina_runs, yield_chart)
        _plot_parameters(parameters, parameter_chart)

        doc = Document()
        _configure_document(doc)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        _font(p.add_run("ПРЕДВАРИТЕЛЬНЫЙ ТЕХНИЧЕСКИЙ ОТЧЁТ"), 9, bold=True, color=BLUE)
        title = doc.add_paragraph()
        title.paragraph_format.space_after = Pt(5)
        _font(title.add_run("Калибровка сортов яровой пшеницы Zlata и Irgina в APSIMX"), 23, bold=True, color=INK)
        subtitle = doc.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(12)
        _font(subtitle.add_run("Использованные данные, метод, полученные параметры и ограничения"), 13, color=MUTED)

        _add_callout(
            doc,
            "Текущий статус.",
            "Выполнено 17 контрольных прогонов. Ошибок APSIMX нет. Получена предварительная калибровка фенологии; урожайность пока не откалибрована.",
        )

        _add_heading(doc, "1. Цель работы")
        _add_body(
            doc,
            "Настроить сорта Zlata и Irgina в APSIMX по сроку колошения. Рассчитанная зрелость не должна наступать позже фактической уборки."
        )

        _add_heading(doc, "2. Исходные данные")
        _add_table(
            doc,
            ["Сорт", "Годы", "Калибровка", "Независимая проверка", "Число сезонов"],
            [
                ["Zlata", "2020–2023", "2020–2022", "2023", str(len(zlata))],
                ["Irgina", "2007–2019", "2007–2016", "2017–2019", str(len(irgina))],
            ],
            [1300, 1750, 1850, 2760, 1700],
        )
        _add_body(
            doc,
            "Источник: «20260804_ВНИИМЗ_трансекта.xlsx», лист «Фенофазы», вариант N30. В расчёт передавались дата посева, суточная погода Open-Meteo и фактическая густота. Глубина посева — 30 мм, междурядье — 250 мм."
        )
        _add_heading(doc, "3. Метод калибровки")
        _add_body(
            doc,
            "Использован grid search: APSIMX рассчитывал каждый набор сортовых параметров по всем годам, после чего выбирался набор с наименьшей ошибкой даты колошения."
        )
        steps = [
            ("1", "Колошение", "Минимум RMSE между расчётной и фактической датой."),
            ("2", "Созревание", "Зрелость не позже даты уборки."),
            ("3", "Проверка", "Контроль на годах, не участвовавших в подборе."),
        ]
        _add_table(doc, ["Этап", "Операция", "Критерий"], [list(x) for x in steps], [650, 2600, 6110], font_size=9.1)
        _add_body(doc, "Критерий: RMSE = √(средняя квадратичная ошибка даты колошения).")
        doc.add_picture(str(phenology_chart), width=Inches(6.55))
        _add_caption(doc, "Рисунок 1. Фактическая и рассчитанная APSIMX дата начала колошения")

        _add_heading(doc, "4. Результаты проверки фенологии")
        metric_rows = []
        for cultivar, summary in (("Zlata", zlata_summary), ("Irgina", irgina_summary)):
            for role, role_label in (("calibration", "Калибровка"), ("validation", "Проверка")):
                n = int(float(summary[f"{role}_n"]))
                maturity_n = int(float(summary[f"{role}_maturity_n"]))
                on_time = int(float(summary[f"{role}_maturity_on_time_count"]))
                metric_rows.append(
                    [
                        cultivar,
                        role_label,
                        str(n),
                        f"{float(summary[f'{role}_heading_mae_days']):.2f}",
                        f"{float(summary[f'{role}_heading_rmse_days']):.2f}",
                        f"{on_time}/{maturity_n}",
                        f"{float(summary[f'{role}_maturity_max_late_days']):.0f}",
                    ]
                )
        metrics_table = _add_table(
            doc,
            ["Сорт", "Выборка", "n", "MAE, дн.", "RMSE, дн.", "Зрелость в срок", "Max опозд., дн."],
            metric_rows,
            [1050, 1500, 500, 1100, 1200, 2100, 1910],
            font_size=8.5,
        )
        _keep_table_together(metrics_table)
        _add_body(
            doc,
            "RMSE на проверке: Irgina — 5,89 дня; Zlata — 7,00 дня. Для Zlata проверка выполнена только по одному году."
        )
        doc.add_picture(str(maturity_chart), width=Inches(6.55))
        _add_caption(doc, "Рисунок 2. Рассчитанная зрелость и фактическая дата уборки")
        _add_callout(
            doc,
            "Проблемные сезоны.",
            "Zlata, 2021: колошение позже на 13 дней, зрелость позже уборки на 18 дней. Irgina, 2019: зрелость позже уборки на 4 дня.",
            fill=CAUTION,
        )

        _add_heading(doc, "5. Полученные сортовые параметры")
        meanings = {
            "MinLN": "Минимальное конечное число листьев",
            "PpLN": "Реакция фенологии на длину дня",
            "BasePhyllochron": "Базовый тепловой интервал появления листьев",
            "GrainFillingTarget": "Тепловая потребность фазы налива зерна",
            "GrainRipeningTarget": "Тепловая потребность фазы созревания",
        }
        parameter_rows = []
        for name in meanings:
            unit = "°C·сут" if name in {"BasePhyllochron", "GrainFillingTarget", "GrainRipeningTarget"} else "коэффициент APSIMX"
            parameter_rows.append([name, meanings[name], unit, f"{parameters['Zlata'][name]:g}", f"{parameters['Irgina'][name]:g}"])
        parameter_table = _add_table(
            doc,
            ["Параметр", "Смысл", "Единица", "Zlata", "Irgina"],
            parameter_rows,
            [1900, 3510, 1900, 1025, 1025],
            font_size=8.6,
        )
        _keep_table_together(parameter_table)
        doc.add_picture(str(parameter_chart), width=Inches(6.55))
        _add_caption(doc, "Рисунок 3. Сравнение текущих сортовых параметров APSIMX")

        _add_heading(doc, "6. Различие сортов")
        _add_body(
            doc,
            "Три из пяти параметров совпали. Irgina отличается меньшими BasePhyllochron и GrainFillingTarget. Для надёжного разделения сортов нужны наблюдения Zlata и Irgina в одинаковые годы и в одинаковых полевых условиях."
        )

        _add_heading(doc, "7. Урожайность: диагностика, а не цель калибровки")
        _add_body(
            doc,
            "Урожайность не входила в критерий калибровки сорта. APSIMX систематически занижает её, поэтому экономические результаты пока следует считать предварительными."
        )
        yield_table = _add_table(
            doc,
            ["Сорт", "Факт, среднее", "APSIMX, среднее", "MAE", "RMSE", "Смещение"],
            [
                ["Zlata", f"{zlata_yield_metrics['observed_mean']:.2f}", f"{zlata_yield_metrics['simulated_mean']:.2f}", f"{zlata_yield_metrics['mae']:.2f}", f"{zlata_yield_metrics['rmse']:.2f}", f"{zlata_yield_metrics['bias']:.2f}"],
                ["Irgina", f"{irgina_yield_metrics['observed_mean']:.2f}", f"{irgina_yield_metrics['simulated_mean']:.2f}", f"{irgina_yield_metrics['mae']:.2f}", f"{irgina_yield_metrics['rmse']:.2f}", f"{irgina_yield_metrics['bias']:.2f}"],
            ],
            [1200, 1900, 2100, 1200, 1200, 1760],
            font_size=8.8,
        )
        _keep_table_together(yield_table)
        doc.add_picture(str(yield_chart), width=Inches(6.55))
        _add_caption(doc, "Рисунок 4. Фактическая и рассчитанная APSIMX урожайность, т/га")
        _add_callout(
            doc,
            "Ограничение.",
            "Zlata: около 0,004 т/га в 2021 году и 0 т/га в 2023 году. Требуется проверка воды, азота и удобрения N30.",
            fill=CAUTION,
        )

        _add_heading(doc, "8. Экономика оптимизации")
        _add_body(
            doc,
            "Экономика рассчитывается после каждого APSIMX-прогона. Урожайность Y поступает из APSIMX; цены и затраты читаются из data/economics/."
        )
        economy_rows = [
            ["Цена зерна", f"{economic_constants['grain_price']:.0f} руб./т"],
            ["Постоянные затраты", f"{economic_constants['constant_C']:.0f} руб./га"],
            ["Внесение удобрения", f"{economic_constants['application_cost']:.0f} руб./га за проход"],
            ["Очистка", f"{economic_constants['cleaning_cost']:.0f} руб./т"],
            ["Погрузка и перевалка", f"{economic_constants['handling_cost']:.0f} руб./т"],
            ["Финансирование", f"{economic_constants['finance_rate'] * 100:.0f}% годовых, {economic_constants['finance_days']:.0f} дней"],
        ]
        economy_table = _add_table(doc, ["Показатель", "Принято в расчёте"], economy_rows, [4300, 5060], font_size=9.1)
        _keep_table_together(economy_table)

        price_rows = [
            [name, f"{fertilizer_prices[name]:.2f}"]
            for name in ("UreaN", "NH4NO3N", "UAN_N", "NO3N")
        ]
        price_table = _add_table(
            doc,
            ["Тип удобрения APSIMX", "Цена, руб./кг внесённого количества"],
            price_rows,
            [4300, 5060],
            font_size=9.1,
        )
        _keep_table_together(price_table)
        _add_body(
            doc,
            "Расчёт: Cуд = Σ(цена × количество); Cфин = 0,18 × 240 / 365 × Cуд; Cупр = Cуд + 480 × число проходов + Cфин + ΔCсемян."
        )
        _add_callout(
            doc,
            "Прибыль по сценарию.",
            "Profit = 13 450 × Y − 23 000 − Cупр, где Y — урожайность, т/га; Profit — руб./га.",
        )
        _add_body(
            doc,
            f"Поправка стоимости семян пока предварительная: базовая густота {seed_config['baseline_population']:.0f} растений/м², масса 1000 семян {seed_config['thousand_seed_weight_g']:.0f} г, всхожесть {seed_config['germination_fraction'] * 100:.0f}%, чистота {seed_config['purity_fraction'] * 100:.0f}%, полевая приживаемость {seed_config['field_establishment_fraction'] * 100:.0f}%, цена {seed_config['seed_price_per_kg']:.0f} руб./кг."
        )
        _add_body(
            doc,
            "Цели NSGA-II: максимальная средняя прибыль, минимальные управляемые затраты и максимальная прибыль в худшем погодном сценарии. Холодный сценарий в текущем config отключён."
        )

        _add_heading(doc, "9. Что требуется уточнить")
        remaining = [
            ["Фенология", "Точная дата физиологической зрелости, а не только дата уборки"],
            ["Условия опыта", "Начальная влажность и азот по слоям, фактический состав и форма удобрения"],
            ["Урожай", "Масса 1000 зёрен, уборочный индекс и урожайность"],
            ["Экономика", "Актуальные цены удобрений и подтверждение параметров стоимости семян"],
            ["Сравнение сортов", "Параллельные наблюдения Zlata и Irgina в одинаковых годах и полевых условиях"],
        ]
        _add_table(doc, ["Блок", "Необходимое уточнение"], remaining, [2200, 7160], font_size=9.2)

        _add_body(
            doc,
            "Итог: фенология настроена предварительно; урожайность и экономические исходные данные требуют уточнения."
        )

        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
