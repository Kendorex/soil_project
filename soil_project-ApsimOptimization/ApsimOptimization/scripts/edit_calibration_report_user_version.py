from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

import build_cultivar_calibration_report as report_builder


ROOT = Path(__file__).resolve().parents[1]
REPORT_SOURCE = ROOT / "docs" / "Отчет_по_калибровке_сортов_Zlata_Irgina.docx"
REPORT_OUTPUT = ROOT / "docs" / "Отчет_по_калибровке_сортов_Zlata_Irgina_чистая_прибыль.docx"
INK = RGBColor(0x24, 0x34, 0x47)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)


def _format_run(run, *, bold: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(10.5)
    run.font.color.rgb = INK
    run.bold = bold


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _math_run(paragraph, text: str, *, bold: bool = False, subscript: bool = False):
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Cambria Math")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Cambria Math")
    run.font.size = Pt(11)
    run.font.color.rgb = INK
    run.bold = bold
    run.font.subscript = subscript
    return run


def _symbol(paragraph, base: str, index: str):
    _math_run(paragraph, base, bold=True)
    return _math_run(paragraph, index, bold=True, subscript=True)


def _replace_economic_formulas(document: Document) -> None:
    calculation = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.startswith("Расчёт:") or paragraph.text.startswith("Cуд =")
    )
    _clear_paragraph(calculation)
    calculation.paragraph_format.space_before = Pt(3)
    calculation.paragraph_format.space_after = Pt(5)

    _symbol(calculation, "C", "уд")
    _math_run(calculation, " = Σ (p")
    _math_run(calculation, "i", subscript=True)
    _math_run(calculation, " × q")
    _math_run(calculation, "i", subscript=True)
    _math_run(calculation, ")").add_break()

    _symbol(calculation, "C", "внес")
    line = _math_run(calculation, " = 480 × n")
    _math_run(calculation, "проходов", subscript=True).add_break()

    _symbol(calculation, "C", "фин")
    _math_run(calculation, " = ")
    _symbol(calculation, "C", "уд")
    line = _math_run(calculation, " × 0,18 × 240 / 365")
    line.add_break()

    _symbol(calculation, "C", "упр")
    _math_run(calculation, " = ")
    _symbol(calculation, "C", "уд")
    _math_run(calculation, " + ")
    _symbol(calculation, "C", "внес")
    _math_run(calculation, " + ")
    _symbol(calculation, "C", "фин")
    _math_run(calculation, " + Δ")
    _symbol(calculation, "C", "сем")
    _math_run(calculation, ".")

    definitions = calculation.add_run(
        "\nгде pᵢ — цена удобрения; qᵢ — внесённое количество; "
        "n — число проходов; ΔCсем — поправка стоимости семян."
    )
    _format_run(definitions)

    profit_cell = next(
        table.cell(0, 0) for table in document.tables
        if table.rows and table.cell(0, 0).text.startswith("Прибыль по сценарию")
    )
    profit = profit_cell.paragraphs[0]
    _clear_paragraph(profit)
    profit.paragraph_format.space_before = Pt(1)
    profit.paragraph_format.space_after = Pt(1)

    heading = profit.add_run("Прибыль по сценарию")
    _format_run(heading, bold=True)
    heading.font.color.rgb = DARK_BLUE
    heading.add_break()

    _symbol(profit, "R", "s")
    _math_run(profit, " = 14 300 × ")
    _symbol(profit, "Y", "s").add_break()

    _math_run(profit, "V", bold=True)
    _math_run(profit, " = (500 + 350) × ")
    _symbol(profit, "Y", "s")
    _math_run(profit, " = 850 × ")
    _symbol(profit, "Y", "s").add_break()

    _symbol(profit, "Π", "s")
    _math_run(profit, " = ")
    _symbol(profit, "R", "s")
    _math_run(profit, " − 23 000 − ")
    _symbol(profit, "C", "упр")
    _math_run(profit, " − ")
    _math_run(profit, "V", bold=True).add_break()

    _symbol(profit, "Π", "s")
    _math_run(profit, " = 13 450 × ")
    _symbol(profit, "Y", "s")
    _math_run(profit, " − 23 000 − ")
    _symbol(profit, "C", "упр")
    _math_run(profit, ".")

    _math_run(profit, "\nИтоговая формула чистой прибыли", bold=True).add_break()
    _math_run(profit, "ЧП", bold=True)
    _math_run(profit, " = 13 450 × Y − 23 000 − ")
    _symbol(profit, "C", "уд")
    _math_run(profit, " − ")
    _symbol(profit, "C", "внес")
    _math_run(profit, " − ")
    _symbol(profit, "C", "фин")
    _math_run(profit, " − Δ")
    _symbol(profit, "C", "сем")
    _math_run(profit, ".")

    explanation = profit.add_run(
        "\nЧП — чистая прибыль, руб./га; Y — урожайность выбранного сценария, т/га; "
        "Yₛ — урожайность в сценарии, т/га; Πₛ — прибыль, руб./га; "
        "Rₛ — выручка; V — затраты на очистку, погрузку и перевалку урожая в рассматриваемом сценарии."
    )
    _format_run(explanation)


def _restore_maturity_chart(document: Document, image_path: Path) -> None:
    caption_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.startswith("Рисунок 2. Рассчитанная зрелость")
    )
    if caption_index == 0:
        raise RuntimeError("Cannot insert maturity chart before its caption")
    image_paragraph = document.paragraphs[caption_index - 1]
    if image_paragraph._p.xpath(".//w:drawing"):
        return
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(2)
    image_paragraph.paragraph_format.space_after = Pt(1)
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.55))


def _mark_preliminary_irgina_value(document: Document) -> None:
    parameter_table = next(
        table for table in document.tables
        if table.rows and table.cell(0, 0).text == "Параметр"
    )
    target_row = next(
        row for row in parameter_table.rows
        if row.cells[0].text == "GrainRipeningTarget"
    )
    irgina_cell = target_row.cells[-1]
    paragraph = irgina_cell.paragraphs[0]
    _clear_paragraph(paragraph)
    value = paragraph.add_run("80*")
    _format_run(value, bold=True)

    note_element = parameter_table._tbl.getnext()
    if note_element is not None and note_element.tag == qn("w:p"):
        possible_note = Paragraph(note_element, parameter_table._parent)
        if possible_note.text.startswith("* Предварительное значение Irgina"):
            return

    note_element = OxmlElement("w:p")
    parameter_table._tbl.addnext(note_element)
    note = Paragraph(note_element, parameter_table._parent)
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(5)
    lead = note.add_run("* Предварительное значение Irgina. ")
    _format_run(lead, bold=True)
    explanation = note.add_run(
        "В быстром тесте 2017–2019 оно сократило максимальное опоздание зрелости с 4 до 3 дней; требуется проверка на полной выборке."
    )
    _format_run(explanation)


def main() -> None:
    document = Document(REPORT_SOURCE)
    zlata_runs, _ = report_builder._load_final_results("zlata")
    irgina_runs, _ = report_builder._load_final_results("irgina")

    with tempfile.TemporaryDirectory(prefix="apsim_report_edit_") as temporary:
        chart = Path(temporary) / "maturity_comparison.png"
        report_builder._plot_maturity_comparison(zlata_runs, irgina_runs, chart)
        _restore_maturity_chart(document, chart)
        _replace_economic_formulas(document)
        document.save(REPORT_OUTPUT)

    print(REPORT_OUTPUT)


if __name__ == "__main__":
    main()
