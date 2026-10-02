#!/usr/bin/env python3
"""Extract reproducible 2019-2023 field observations from the source workbook."""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.multifield_model import inventory, write_json
from src.multifield_optimizer import ROOT, load_config


DEFAULT_WORKBOOK = ROOT.parent/'more_calibration'/'20260804_ВНИИМЗ_трансекта ЧИСТАЯ.xlsx'
DEFAULT_OUTPUT = ROOT/'data'/'observations'/'multifield_2019_2023.json'


def present(value):
    return value is not None and str(value).strip() not in {'', 'Нет данных'}


def crop_group(value):
    text = str(value).strip().lower()
    if 'овес' in text or 'овёс' in text:
        return 'Овёс'
    if 'клеверотимофе' in text or 'многолетн' in text:
        return 'Клевер + тимофеевка'
    if 'пшениц' in text:
        return 'Яровая пшеница'
    if 'рож' in text:
        return 'Озимая рожь'
    if 'картоф' in text:
        return 'Картофель'
    if 'рапс' in text:
        return 'Рапс'
    if 'горч' in text:
        return 'Горчица'
    return str(value).strip()


def parse_day_month(value, harvest_year, winter=False):
    if not present(value):
        return None
    year = harvest_year - 1 if winter else harvest_year
    if isinstance(value, (dt.datetime, dt.date)):
        return dt.date(year, value.month, value.day).isoformat()
    text = str(value).strip().replace(',', '.')
    match = re.fullmatch(r'(\d{1,2})\.(\d{1,2})(?:\.\d{2,4})?', text)
    if not match:
        return None
    day, month = map(int, match.groups())
    try:
        return dt.date(year, month, day).isoformat()
    except ValueError:
        return None


def components(group):
    return {
        'Овёс': ['Oats'],
        'Яровая пшеница': ['Wheat'],
        'Озимая рожь': ['Wheat'],
        'Клевер + тимофеевка': ['RedClover', 'PlantainForage'],
        'Картофель': ['Potato'],
        'Рапс': ['Rape'],
        'Горчица': ['Mustard'],
    }.get(group, [])


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument('--config', default=str(ROOT/'configs'/'multifield.yaml'))
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)

    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        raise SystemExit(f'Missing workbook: {workbook_path}')
    config = load_config(args.config)
    meta = inventory(config)
    book = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if 'Фенофазы' not in book.sheetnames:
        raise ValueError('Workbook has no Фенофазы sheet')
    sheet = book['Фенофазы']
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    required = {'Год', 'Поле', 'Культура', 'Сорт', 'Технология',
                'Средняя урожайность, т/га', 'Дата посева', 'Дата уборки'}
    if required - set(headers):
        raise ValueError(f'Missing workbook columns: {sorted(required-set(headers))}')

    years = {}
    corrections = []
    for values in rows:
        raw = dict(zip(headers, values))
        year, number = raw.get('Год'), raw.get('Поле')
        observed_yield = raw.get('Средняя урожайность, т/га')
        if (not isinstance(year, (int, float)) or not isinstance(number, (int, float)) or
                not isinstance(observed_yield, (int, float))):
            continue
        year, number = int(year), int(number)
        if not 2019 <= year <= 2023 or not 1 <= number <= 10:
            continue
        if not math.isfinite(float(observed_yield)) or observed_yield < 0:
            raise ValueError(f'Invalid observed yield for {year}/Field{number}')
        field = f'Field{number}'
        index = year - int(meta['start'][:4])
        rotation = meta['fields'][field]['rotation']
        model_crop = rotation['RotationCrops'].split(',')[index].strip()
        model_cultivar = rotation['RotationCultivars'].split(',')[index].strip()
        group = crop_group(raw.get('Культура'))
        winter = group == 'Озимая рожь'
        record = {
            'field': field,
            'crop': str(raw.get('Культура') or '').strip(),
            'crop_group': group,
            'cultivar': re.sub(r'\s+', ' ', str(raw.get('Сорт') or '').strip()),
            'technology': str(raw.get('Технология') or '').strip(),
            'observed_yield_t_ha': float(observed_yield),
            'observed_sow_date': parse_day_month(raw.get('Дата посева'), year, winter),
            'observed_harvest_date': parse_day_month(raw.get('Дата уборки'), year),
            'observed_tillering_date': parse_day_month(
                raw.get('Фенофаза кущения, дата начала'), year, winter),
            'observed_heading_date': parse_day_month(
                raw.get('Фенофаза колошения, дата начала'), year),
            'model_crop': model_crop,
            'model_cultivar': model_cultivar,
            'yield_components': components(group),
        }
        # Confirmed correction retained from the source-data audit.
        if year == 2023 and field == 'Field2' and group == 'Овёс':
            original = record['observed_tillering_date']
            record['observed_tillering_date'] = '2023-06-05'
            corrections.append({'year': year, 'field': field,
                                'column': 'observed_tillering_date',
                                'source_value': original, 'corrected_value': '2023-06-05',
                                'reason': 'confirmed source typo: 05.05 -> 05.06'})
        years.setdefault(str(year), []).append(record)

    for year, records in years.items():
        records.sort(key=lambda row: int(row['field'][5:]))
        seen = [row['field'] for row in records]
        if len(seen) != len(set(seen)):
            raise ValueError(f'Duplicate field observation in {year}')
        for row in records:
            if not row['yield_components']:
                raise ValueError(f'Unsupported observed crop group: {row["crop_group"]}')

    payload = {
        'schema_version': 1,
        'source': {
            'workbook': str(workbook_path),
            'workbook_sha256': hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
            'sheet': 'Фенофазы',
            'yield_column': 'Средняя урожайность, т/га',
            'extracted_utc': dt.datetime.now(dt.timezone.utc).isoformat(),
        },
        'model_mapping': {
            'model_sha256': meta['model_sha256'],
            'clock_start': meta['start'],
            'clock_end': meta['end'],
        },
        'source_corrections': corrections,
        'complete_field_coverage_by_year': {
            year: len(records) == len(meta['fields']) for year, records in years.items()},
        'interpretation_limits': [
            'The workbook states t/ha but does not identify dry, fresh, or marketable mass for every crop.',
            'Missing harvest dates cannot be used to validate APSIM harvest timing.',
            'Technology labels do not establish fertilizer product, nutrient fraction, application date, or depth.',
            'PlantainForage is the model proxy for timothy and is a structural model approximation.',
        ],
        'years': dict(sorted(years.items())),
    }
    if sum(map(len, payload['years'].values())) != 37:
        raise ValueError(f'Expected 37 observations, got {sum(map(len, payload["years"].values()))}')
    write_json(args.output, payload)
    print(f'OBSERVATIONS: {args.output.resolve()}')
    print('Records: ' + ', '.join(f'{year}={len(rows)}' for year, rows in payload['years'].items()))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
