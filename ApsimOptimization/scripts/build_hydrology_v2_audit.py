from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import fmean
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_APSIMX = ROOT / "simulation" / "standart.apsimx"
DEFAULT_WORKBOOK = (
    ROOT
    / "docs"
    / "20260804_ВНИИМЗ_трансекта"
    / "20260804_ВНИИМЗ_трансекта.xlsx"
)
DEFAULT_OUTPUT_DIR = ROOT / "experiments" / "hydrology_v2" / "audit"


class AuditError(RuntimeError):
    pass


def _named(node: Any, name: str) -> dict[str, Any]:
    if isinstance(node, dict):
        if node.get("Name") == name:
            return node
        for value in node.values():
            try:
                return _named(value, name)
            except AuditError:
                pass
    elif isinstance(node, list):
        for value in node:
            try:
                return _named(value, name)
            except AuditError:
                pass
    raise AuditError(f"APSIMX node {name!r} was not found")


def _parse_day_month(value: Any, year: int) -> date:
    if isinstance(value, datetime):
        return value.date().replace(year=year)
    if isinstance(value, date):
        return value.replace(year=year)
    text = str(value).strip().replace(",", ".")
    parts = text.split(".")
    if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
        return date(year, int(parts[1]), int(parts[0]))
    raise AuditError(f"Cannot parse day/month value {value!r} for {year}")


def _parse_groundwater_period(value: Any, year: int) -> tuple[date, str, str]:
    try:
        parsed = _parse_day_month(value, year)
        return parsed, "day", "source_day_month"
    except AuditError:
        pass
    month_names = {
        "январь": 1,
        "февраль": 2,
        "март": 3,
        "апрель": 4,
        "май": 5,
        "июнь": 6,
        "июль": 7,
        "август": 8,
        "сентябрь": 9,
        "сентбярь": 9,
        "октябрь": 10,
        "ноябрь": 11,
        "декабрь": 12,
    }
    text = str(value).strip().lower()
    month = month_names.get(text)
    if month is None:
        raise AuditError(f"Cannot parse groundwater period {value!r} for {year}")
    return date(year, month, 15), "month", "month_midpoint_day_15_due_unknown_source_day"


def _as_float(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _profile_rows(apsimx_path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    data = json.loads(apsimx_path.read_text(encoding="utf-8"))
    soil = _named(data, "Soil")
    physical = _named(data, "Physical")
    water = _named(data, "Water")
    crop = _named(data, "WheatSoil")
    arrays = {
        "thickness_mm": list(physical.get("Thickness") or []),
        "bd_g_cm3": list(physical.get("BD") or []),
        "airdry_m3_m3": list(physical.get("AirDry") or []),
        "ll15_m3_m3": list(physical.get("LL15") or []),
        "dul_m3_m3": list(physical.get("DUL") or []),
        "sat_m3_m3": list(physical.get("SAT") or []),
        "initial_water_m3_m3": list(water.get("InitialValues") or []),
        "wheat_ll_m3_m3": list(crop.get("LL") or []),
        "wheat_kl_per_day": list(crop.get("KL") or []),
        "wheat_xf": list(crop.get("XF") or []),
    }
    layer_count = len(arrays["thickness_mm"])
    if not layer_count or any(len(values) != layer_count for values in arrays.values()):
        raise AuditError("Soil profile arrays do not have equal non-zero lengths")

    rows = []
    top = 0.0
    for index in range(layer_count):
        thickness = float(arrays["thickness_mm"][index])
        row = {
            "layer": index + 1,
            "top_mm": top,
            "bottom_mm": top + thickness,
        }
        row.update({name: values[index] for name, values in arrays.items()})
        row["ordering_valid"] = (
            float(row["airdry_m3_m3"])
            <= float(row["ll15_m3_m3"])
            < float(row["dul_m3_m3"])
            < float(row["sat_m3_m3"])
        )
        rows.append(row)
        top += thickness

    metadata = {
        "apsimx": apsimx_path.relative_to(ROOT).as_posix(),
        "apsimx_sha256": _sha256(apsimx_path),
        "soil_data_source": soil.get("DataSource"),
        "soil_comments": soil.get("Comments"),
        "soil_latitude": soil.get("Latitude"),
        "soil_longitude": soil.get("Longitude"),
        "layer_count": layer_count,
    }
    return rows, metadata


def _load_field_observations(
    workbook_path: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        density_rows = list(workbook["Плотность"].iter_rows(values_only=True))[1:]
        moisture_rows = list(workbook["Влажность"].iter_rows(values_only=True))[1:]
        groundwater_rows = list(workbook["УГВ"].iter_rows(values_only=True))[1:]

        density_by_point: dict[tuple[int, int, int], float] = {}
        density_depth_by_point: dict[tuple[int, int, int], str] = {}
        density_raw = []
        for row in density_rows:
            try:
                key = (int(row[0]), int(row[1]), int(row[2]))
            except (TypeError, ValueError):
                continue
            density = _as_float(row[3])
            if density is None:
                continue
            density_by_point[key] = density
            density_depth_by_point[key] = str(row[4]).strip()
            density_raw.append(
                {
                    "year": key[0],
                    "field": key[1],
                    "point": key[2],
                    "bulk_density_g_cm3": density,
                    "density_depth": density_depth_by_point[key],
                    "source_sheet": "Плотность",
                }
            )

        paired = []
        for row in moisture_rows:
            try:
                year, field, point = int(row[0]), int(row[1]), int(row[2])
                measured_on = _parse_day_month(row[3], year)
            except (TypeError, ValueError, AuditError):
                continue
            gravimetric_percent = _as_float(row[4])
            density = density_by_point.get((year, field, point))
            if gravimetric_percent is None or density is None:
                continue
            volumetric = gravimetric_percent / 100.0 * density
            porosity = 1.0 - density / 2.65
            if gravimetric_percent < 0.0:
                quality_flag = "invalid_negative_moisture"
            elif not 0.0 < density < 2.65:
                quality_flag = "invalid_bulk_density"
            elif volumetric > porosity:
                quality_flag = "invalid_above_particle_density_porosity"
            else:
                quality_flag = "ok"
            paired.append(
                {
                    "year": year,
                    "field": field,
                    "point": point,
                    "measurement_date": measured_on.isoformat(),
                    "measurement_depth": "0-20 cm",
                    "gravimetric_water_percent": gravimetric_percent,
                    "bulk_density_g_cm3": density,
                    "volumetric_water_m3_m3": volumetric,
                    "porosity_from_density_m3_m3": porosity,
                    "quality_flag": quality_flag,
                    "conversion": "theta_v = theta_g_percent / 100 * bulk_density_g_cm3",
                    "source_sheets": "Влажность + Плотность",
                }
            )

        groundwater = []
        for row in groundwater_rows:
            try:
                year = int(row[0])
                measured_on, date_precision, date_assumption = _parse_groundwater_period(row[4], year)
            except (TypeError, ValueError, AuditError):
                continue
            depth_m = _as_float(row[3])
            if depth_m is None:
                continue
            groundwater.append(
                {
                    "year": year,
                    "landscape": row[1],
                    "points": row[2],
                    "measurement_date": measured_on.isoformat(),
                    "measurement_period": measured_on.strftime("%Y-%m"),
                    "date_precision": date_precision,
                    "date_assumption": date_assumption,
                    "groundwater_depth_m": depth_m,
                    "groundwater_depth_mm": depth_m * 1000.0,
                    "source_sheet": "УГВ",
                }
            )
        return density_raw, paired, groundwater
    finally:
        workbook.close()


def _summarise_topsoil(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[int, int, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[(int(row["year"]), int(row["field"]), str(row["measurement_date"]))].append(row)
    summary = []
    for (year, field, measured_on), items in sorted(grouped.items()):
        valid = [item for item in items if item["quality_flag"] == "ok"]
        if not valid:
            continue
        values = [float(item["volumetric_water_m3_m3"]) for item in valid]
        densities = [float(item["bulk_density_g_cm3"]) for item in valid]
        summary.append(
            {
                "year": year,
                "field": field,
                "measurement_date": measured_on,
                "measurement_depth": "0-20 cm",
                "point_count_raw": len(items),
                "point_count_valid": len(valid),
                "point_count_excluded": len(items) - len(valid),
                "bulk_density_mean_g_cm3": fmean(densities),
                "volumetric_water_mean_m3_m3": fmean(values),
                "volumetric_water_min_m3_m3": min(values),
                "volumetric_water_max_m3_m3": max(values),
            }
        )
    return summary


def build_audit(apsimx_path: Path, workbook_path: Path, output_dir: Path) -> None:
    profile, metadata = _profile_rows(apsimx_path)
    density, topsoil, groundwater = _load_field_observations(workbook_path)
    topsoil_summary = _summarise_topsoil(topsoil)

    model_bd_top = float(profile[0]["bd_g_cm3"])
    observed_density = [float(row["bulk_density_g_cm3"]) for row in density]
    observed_density_mean = fmean(observed_density)
    field_2023_density = [
        float(row["bulk_density_g_cm3"])
        for row in density
        if int(row["year"]) == 2023 and int(row["field"]) == 3
    ]
    recent = [row for row in topsoil_summary if int(row["year"]) >= 2020]
    model_ll15_top = float(profile[0]["ll15_m3_m3"])

    audit = {
        **metadata,
        "workbook": workbook_path.relative_to(ROOT).as_posix(),
        "workbook_sha256": _sha256(workbook_path),
        "model_top_layer_bd_g_cm3": model_bd_top,
        "observed_density_mean_g_cm3": observed_density_mean,
        "observed_density_count": len(observed_density),
        "observed_2023_field3_density_mean_g_cm3": fmean(field_2023_density),
        "observed_2023_field3_density_count": len(field_2023_density),
        "model_top_layer_ll15_m3_m3": model_ll15_top,
        "recent_topsoil_observation_groups": len(recent),
        "recent_groups_below_model_top_ll15": sum(
            float(row["volumetric_water_mean_m3_m3"]) < model_ll15_top for row in recent
        ),
        "topsoil_measurements_excluded_by_quality_flag": sum(
            row["quality_flag"] != "ok" for row in topsoil
        ),
        "groundwater_observation_count": len(groundwater),
        "groundwater_years": sorted({int(row["year"]) for row in groundwater}),
        "groundwater_2023_landscapes": sorted(
            {int(row["landscape"]) for row in groundwater if int(row["year"]) == 2023}
        ),
        "confirmed_findings": [
            "The active APSIMX soil profile is sourced from an Australian CSIRO reference soil, not this field.",
            "Measured bulk density differs materially from the active profile value.",
            "Some recent observed 0-20 cm water contents are below the active top-layer LL15.",
        ],
        "not_identified_from_available_measurements": [
            "layer-specific AirDry",
            "layer-specific LL15",
            "layer-specific DUL",
            "layer-specific SAT",
            "layer-specific Ks and retention-curve parameters",
            "continuous groundwater boundary for years other than the measured period",
        ],
        "status": "audit_only_no_soil_parameters_changed",
    }

    _write_csv(output_dir / "model_soil_profile.csv", profile)
    _write_csv(output_dir / "bulk_density_observations.csv", density)
    _write_csv(output_dir / "topsoil_water_observations.csv", topsoil)
    _write_csv(output_dir / "topsoil_water_summary.csv", topsoil_summary)
    _write_csv(output_dir / "groundwater_observations.csv", groundwater)
    (output_dir / "audit_summary.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"HYDROLOGY V2 AUDIT OK: {output_dir}")
    print(f"Model soil source: {metadata['soil_data_source']}")
    print(f"Model top BD={model_bd_top:.3f}; observed mean BD={observed_density_mean:.3f}")
    print(
        "Recent topsoil groups below model LL15: "
        f"{audit['recent_groups_below_model_top_ll15']}/{audit['recent_topsoil_observation_groups']}"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a read-only hydrology-v2 input audit")
    parser.add_argument("--apsimx", type=Path, default=DEFAULT_APSIMX)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    build_audit(args.apsimx.resolve(), args.workbook.resolve(), args.output_dir.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
