#!/usr/bin/env python3
"""Screen data-derived oat grain-number settings after the residue correction."""
from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import hashlib
import io
import json
import tarfile
import sys
from pathlib import Path
from statistics import fmean

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.cultivar_yield_calibration import upsert_named_cultivar_commands
from src.historical_case import load_observations, records_for_year
from src.multifield_model import fingerprint, inventory, prepare, read_json, write_json
from src.multifield_optimizer import ROOT, load_config, run_apsim
from src.oats_phenology import phenology_candidate_metrics
from src.oats_physiology import collect_oats_diagnostics, patch_oats_diagnostic_reports
from src.oats_residue_sensitivity import (
    defer_wheat_auto_harvest_for_field_years,
    patch_export_new_residue_at_cereal_to_oats_harvest,
)


YEARS = tuple(range(2019, 2024))
FIELDS_2023 = ("Field2", "Field7")
GRAIN_NUMBER = "[Grain].NumberFunction.GrainNumber.GrainsPerGramOfStem.FixedValue"
CURRENT_GRAINS_PER_GRAM_STEM = 45.0
DEFAULT_TRANSITION_RESULTS = (
    ROOT / "outputs" / "multifield" / "OATS_WHEAT_RESIDUE"
    / "20260929_120748" / "documented_2022_harvest_export_new_residue"
    / "oats_2023.csv"
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def observed_grain_sizes(observations: dict) -> dict[str, float]:
    workbook_path = Path(observations["source"]["workbook"])
    if fingerprint(workbook_path) != observations["source"].get("workbook_sha256"):
        raise ValueError("source workbook hash does not match observations")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Фенофазы"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {str(name): position for position, name in enumerate(header)}
    result = {}
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        if cells[index["Год"]] != 2023 or cells[index["Поле"]] not in (2, 7):
            continue
        crop = str(cells[index["Культура"]]).casefold().replace("ё", "е")
        if "овес" not in crop:
            continue
        mass = cells[index["Масса 1000 зёрен, г"]]
        if not isinstance(mass, (int, float)):
            raise ValueError("missing observed thousand-grain mass")
        result[f"Field{int(cells[index['Поле']])}"] = float(mass) / 1000.0
    workbook.close()
    if set(result) != set(FIELDS_2023):
        raise ValueError(f"missing 2023 grain-size observations: {result}")
    return result


def derive_profiles(observations: dict, transition_results: Path):
    observed_sizes = observed_grain_sizes(observations)
    rows = {row["field"]: row for row in read_csv(transition_results)}
    derivation = []
    derived_values = []
    for field in FIELDS_2023:
        row = rows[field]
        observed_yield_g_m2 = float(row["observed_yield_t_ha"]) * 100.0
        implied_number = observed_yield_g_m2 / observed_sizes[field]
        model_number = float(row["final_grain_number_m2"])
        value = CURRENT_GRAINS_PER_GRAM_STEM * implied_number / model_number
        derived_values.append(value)
        derivation.append({
            "field": field,
            "observed_grain_size_g": observed_sizes[field],
            "observed_yield_t_ha": float(row["observed_yield_t_ha"]),
            "observed_yield_implied_grain_number_m2": implied_number,
            "transition_model_grain_number_m2": model_number,
            "current_grains_per_gram_stem": CURRENT_GRAINS_PER_GRAM_STEM,
            "data_derived_grains_per_gram_stem_if_linear": value,
        })
    low, high = sorted(derived_values)
    midpoint = fmean(derived_values)
    profiles = (
        {"name": "transition_control_45", "value": CURRENT_GRAINS_PER_GRAM_STEM,
         "source": "current Yakov setting"},
        {"name": "source_field7_derived", "value": low,
         "source": "Field7 observed yield and thousand-grain mass"},
        {"name": "source_mean_derived", "value": midpoint,
         "source": "mean of Field2 and Field7 data-derived values"},
        {"name": "source_field2_derived", "value": high,
         "source": "Field2 observed yield and thousand-grain mass"},
    )
    return profiles, derivation, observed_sizes


def run_profile(config, observations, year_records, oat_fields, profile, output):
    output.mkdir(parents=True, exist_ok=False)
    model = read_json(config["model"])
    lifecycle_changes = defer_wheat_auto_harvest_for_field_years(
        model, {field: {2022} for field in FIELDS_2023}
    )
    residue_changes = patch_export_new_residue_at_cereal_to_oats_harvest(
        model, FIELDS_2023, {field: {2022} for field in FIELDS_2023}
    )
    cultivar_changes = upsert_named_cultivar_commands(
        model, "Yakov", {GRAIN_NUMBER: profile["value"]}, expected_count=10
    )
    report_changes = patch_oats_diagnostic_reports(model, oat_fields)
    source = output / "source_candidate.apsimx"
    write_json(source, model)
    case = copy.deepcopy(config)
    case["model"] = str(source.resolve())
    case["model_provenance"] = None
    case["historical_replay"] = dict(config["historical_replay"], enabled=True)
    case["scenario"] = {
        "kind": "diagnostic", "name": profile["name"],
        "recommendation_status": "prohibited",
    }
    meta = inventory(case)
    write_json(output / "config_snapshot.json", case)
    write_json(output / "model_inventory.json", meta)
    prepared = prepare(case, output)
    database = run_apsim(case, prepared)

    rows, schemas = [], {}
    for year in YEARS:
        _, summaries, schema = collect_oats_diagnostics(
            database, meta, year_records[year], year=year,
            scenario=profile["name"],
        )
        rows.extend(summaries)
        schemas[str(year)] = schema
    metrics = phenology_candidate_metrics(rows)
    write_csv(output / "oats_history.csv", rows)
    write_json(output / "report_schema.json", schemas)
    result = {
        **profile, "metrics": metrics,
        "lifecycle_changes": lifecycle_changes,
        "residue_rotation_manager_paths": residue_changes,
        "cultivar_changes": cultivar_changes,
        "report_changes": report_changes,
    }
    write_json(output / "profile_summary.json", result)
    return result, rows


def create_archive(path: Path, root: Path) -> None:
    allowed = {
        "summary.json", "candidate_metrics.csv", "case_predictions.csv",
        "grain_component_comparison.csv", "parameter_derivation.csv",
        "profile_summary.json", "oats_history.csv", "report_schema.json",
        "config_snapshot.json", "model_inventory.json", "preparation.json", "apsim.log",
    }
    files = sorted(
        file for file in root.rglob("*") if file.is_file() and file.name in allowed
    )
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [{
            "path": file.relative_to(root).as_posix(),
            "size": file.stat().st_size,
            "sha256": hashlib.sha256(file.read_bytes()).hexdigest(),
        } for file in files],
    }
    with tarfile.open(path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.relative_to(root).as_posix(), recursive=False)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", action="store_true", required=True)
    parser.add_argument(
        "--config", type=Path,
        default=ROOT / "configs" / "multifield_historical_replay.yaml",
    )
    parser.add_argument(
        "--transition-results", type=Path, default=DEFAULT_TRANSITION_RESULTS,
    )
    args = parser.parse_args(argv)

    config = load_config(args.config)
    observations = load_observations(config["historical_observations"])
    profiles, derivation, observed_sizes = derive_profiles(
        observations, args.transition_results
    )
    year_records = {year: records_for_year(observations, year) for year in YEARS}
    oat_fields = tuple(dict.fromkeys(
        row["field"] for year in YEARS for row in year_records[year]
        if "Oats" in row.get("yield_components", [])
    ))
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    root = ROOT / "outputs" / "multifield" / "OATS_GRAIN_NUMBER_SOURCE" / stamp
    root.mkdir(parents=True, exist_ok=False)
    write_csv(root / "parameter_derivation.csv", derivation)

    results, case_rows, component_rows = [], [], []
    for index, profile in enumerate(profiles, 1):
        print(f"[{index}/{len(profiles)}] {profile['name']}", flush=True)
        result, rows = run_profile(
            config, observations, year_records, oat_fields, profile,
            root / profile["name"],
        )
        results.append(result)
        case_rows.extend({"candidate": profile["name"], **row} for row in rows)
        for row in rows:
            if int(row["year"]) != 2023 or row["field"] not in FIELDS_2023:
                continue
            observed_size = observed_sizes[row["field"]]
            implied_number = float(row["observed_yield_t_ha"]) * 100 / observed_size
            component_rows.append({
                "candidate": profile["name"], "field": row["field"],
                "grains_per_gram_stem": profile["value"],
                "observed_yield_t_ha": row["observed_yield_t_ha"],
                "simulated_yield_t_ha": row["simulated_yield_t_ha"],
                "observed_grain_size_g": observed_size,
                "simulated_grain_size_g": row["final_grain_size_g"],
                "observed_yield_implied_grain_number_m2": implied_number,
                "simulated_grain_number_m2": row["final_grain_number_m2"],
                "simulated_harvest_index": row["simulated_harvest_index"],
                "grain_n_allocation_to_demand_ratio": (
                    row["grain_n_allocation_to_demand_ratio"]
                ),
            })
        dev = result["metrics"]["development_yakov_2020_2022"]
        val = result["metrics"]["validation_yakov_2023"]
        print(
            f"  development RMSE={dev['yield_rmse_t_ha']:.3f}; "
            f"2023 RMSE={val['yield_rmse_t_ha']:.3f}", flush=True,
        )

    metric_rows = []
    for result in results:
        for scope, metrics in result["metrics"].items():
            metric_rows.append({
                "candidate": result["name"], "scope": scope,
                "grains_per_gram_stem": result["value"],
                "yield_mae_t_ha": metrics["yield_mae_t_ha"],
                "yield_rmse_t_ha": metrics["yield_rmse_t_ha"],
                "yield_bias_t_ha": metrics["yield_bias_t_ha"],
                "tillering_mae_days": metrics["tillering"]["mae_days"],
                "heading_mae_days": metrics["heading"]["mae_days"],
            })
    write_csv(root / "candidate_metrics.csv", metric_rows)
    write_csv(root / "case_predictions.csv", case_rows)
    write_csv(root / "grain_component_comparison.csv", component_rows)
    write_json(root / "summary.json", {
        "experiment": "oats_data_derived_grain_number_screen_after_residue_correction",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": "engineering_diagnostic_only", "base_model_changed": False,
        "profiles": profiles, "parameter_derivation": derivation,
        "inputs": {
            "model": config["model"], "model_sha256": fingerprint(config["model"]),
            "weather": config["weather"], "weather_sha256": fingerprint(config["weather"]),
            "observations": config["historical_observations"],
            "observations_sha256": fingerprint(config["historical_observations"]),
            "transition_results": str(args.transition_results),
            "transition_results_sha256": fingerprint(args.transition_results),
        },
        "interpretation_limits": [
            "Candidate values are derived from observed yield and thousand-grain mass, assuming a locally linear response of grain number to the cultivar command.",
            "The 2023 transition scenario assumes complete export of newly created wheat harvest residue; this management is not confirmed.",
            "Observed yield moisture basis is not stated.",
            "Candidates are screened against 2020-2022 development cases and are not applied to the calibrated model automatically.",
        ],
    })
    archive = ROOT / f"oats_grain_number_source_{stamp}.tar.gz"
    create_archive(archive, root)
    print(f"OUTPUT: {root}", flush=True)
    print(f"SEND THIS ARCHIVE: {archive}", flush=True)
    print(f"ARCHIVE SHA256: {fingerprint(archive)}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
