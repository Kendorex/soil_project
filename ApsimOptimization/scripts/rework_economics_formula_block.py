from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

import build_cultivar_calibration_report as report_builder
import edit_calibration_report_user_version as editor
from src.config import load_config, load_inputs


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Отчет_по_калибровке_сортов_Zlata_Irgina_чистая_прибыль.docx"
OUTPUT = ROOT / "docs" / "Отчет_по_калибровке_сортов_Zlata_Irgina_итоговый.docx"


def _remove_block(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _write_formula_block(document: Document) -> object:
    profit_table = next(
        table for table in document.tables
        if table.rows and table.cell(0, 0).text.startswith("Прибыль по сценарию")
    )
    paragraph = profit_table.cell(0, 0).paragraphs[0]
    editor._clear_paragraph(paragraph)
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)

    heading = paragraph.add_run("Основная формула чистой прибыли")
    editor._format_run(heading, bold=True)
    heading.font.color.rgb = editor.DARK_BLUE
    heading.add_break()

    editor._symbol(paragraph, "Π", "s")
    editor._math_run(paragraph, " = ")
    editor._symbol(paragraph, "p", "g")
    editor._math_run(paragraph, " × ")
    editor._symbol(paragraph, "Y", "s")
    editor._math_run(paragraph, " − ")
    editor._symbol(paragraph, "C", "0")
    editor._math_run(paragraph, " − ")
    editor._symbol(paragraph, "C", "тех")
    editor._math_run(paragraph, " − V").add_break()

    label = paragraph.add_run("Расчёт составляющих")
    editor._format_run(label, bold=True)
    label.font.color.rgb = editor.DARK_BLUE
    label.add_break()

    editor._math_run(paragraph, "V", bold=True)
    editor._math_run(paragraph, " = (")
    editor._symbol(paragraph, "c", "h")
    editor._math_run(paragraph, " + ")
    editor._symbol(paragraph, "c", "clean")
    editor._math_run(paragraph, " + ")
    editor._symbol(paragraph, "c", "hand")
    editor._math_run(paragraph, ") × ")
    editor._symbol(paragraph, "Y", "s").add_break()

    editor._symbol(paragraph, "C", "тех")
    editor._math_run(paragraph, " = ")
    editor._symbol(paragraph, "C", "F")
    editor._math_run(paragraph, " + ")
    editor._symbol(paragraph, "C", "app")
    editor._math_run(paragraph, " + ")
    editor._symbol(paragraph, "C", "fin")
    editor._math_run(paragraph, " + Δ")
    editor._symbol(paragraph, "C", "seed").add_break()

    editor._symbol(paragraph, "C", "F")
    editor._math_run(paragraph, " = Σ (")
    editor._symbol(paragraph, "p", "i")
    editor._math_run(paragraph, " × ")
    editor._symbol(paragraph, "q", "i")
    editor._math_run(paragraph, ")").add_break()

    editor._symbol(paragraph, "C", "app")
    editor._math_run(paragraph, " = ")
    editor._symbol(paragraph, "c", "app")
    editor._math_run(paragraph, " × ")
    editor._symbol(paragraph, "n", "pass").add_break()

    editor._symbol(paragraph, "C", "fin")
    editor._math_run(paragraph, " = r × (")
    editor._symbol(paragraph, "τ", "")
    editor._math_run(paragraph, " / ")
    editor._symbol(paragraph, "D", "y")
    editor._math_run(paragraph, ") × ")
    editor._symbol(paragraph, "C", "F")

    tr_properties = profit_table.rows[0]._tr.get_or_add_trPr()
    tr_properties.append(report_builder.OxmlElement("w:cantSplit"))

    return profit_table


def _symbol_rows(config: dict, inputs: dict) -> list[list[str]]:
    constants = inputs["constants"]
    prices = inputs["prices"]
    price_text = "; ".join(
        [
            f"UreaN — {prices['UreaN']:.2f}",
            f"NH4NO3N — {prices['NH4NO3N']:.2f}",
            f"UAN_N — {prices['UAN_N']:.2f}",
            f"NO3N — {prices['NO3N']:.2f}",
        ]
    ).replace(".", ",")
    return [
        ["Πₛ", "Чистая прибыль в сценарии s", "Результат расчёта", "руб./га"],
        ["Yₛ", "Урожайность APSIMX в сценарии s", "Результат APSIMX", "т/га"],
        ["p_g", "Цена зерна", f"{constants['grain_price']:,.0f}".replace(",", " "), "руб./т"],
        ["C₀", "Постоянные затраты", f"{constants['constant_C']:,.0f}".replace(",", " "), "руб./га"],
        ["C_тех", "Затраты выбранной технологии", "Рассчитывается", "руб./га"],
        ["V", "Затраты, зависящие от урожайности", "Рассчитывается", "руб./га"],
        ["c_h; c_clean; c_hand", "Уборка; очистка; погрузка и перевалка", "0; 500; 350", "руб./т"],
        ["C_F", "Стоимость удобрений", "Рассчитывается", "руб./га"],
        ["p_i", "Цена i-го типа удобрения", price_text, "руб./кг количества"],
        ["q_i", "Количество i-го удобрения", "По схеме удобрения", "кг/га"],
        ["C_app", "Стоимость внесения", "Рассчитывается", "руб./га"],
        ["c_app", "Стоимость одного прохода", f"{constants['application_cost']:.0f}", "руб./га за проход"],
        ["n_pass", "Число проходов", "По схеме удобрения", "проход"],
        ["C_fin", "Стоимость финансирования", "Рассчитывается", "руб./га"],
        ["r", "Годовая ставка финансирования", f"{constants['finance_rate']:.2f}".replace(".", ","), "доля/год"],
        ["τ", "Период финансирования", f"{constants['finance_days']:.0f}", "дней"],
        ["D_y", "Число дней в году", "365", "дней"],
        ["ΔC_seed", "Поправка стоимости семян", "Рассчитывается; предварительно", "руб./га"],
    ]


def _replace_economic_table(document: Document, profit_table, rows: list[list[str]]) -> None:
    old_table = next(
        table for table in document.tables
        if table.rows and table.cell(0, 0).text == "Показатель"
    )
    new_table = document.add_table(rows=1, cols=4)
    new_table.style = "Table Grid"
    for index, value in enumerate(["Символ", "Что обозначает", "Значение", "Единица"]):
        new_table.cell(0, index).text = value
    for values in rows:
        cells = new_table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    report_builder._format_table(new_table, header=True, font_size=8.2)
    report_builder._set_table_geometry(new_table, [1350, 3250, 2910, 1850])
    new_table.rows[0]._tr.get_or_add_trPr().append(report_builder.OxmlElement("w:tblHeader"))

    profit_table._tbl.addnext(new_table._tbl)
    _remove_block(old_table._tbl)


def main() -> None:
    document = Document(SOURCE)
    config = load_config(ROOT / "configs" / "optimizer_config.yaml")
    inputs = load_inputs(config)

    profit_table = _write_formula_block(document)
    _replace_economic_table(document, profit_table, _symbol_rows(config, inputs))

    economy_heading = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text == "5. Экономика оптимизации"
    )
    economy_heading.paragraph_format.page_break_before = True

    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("Cуд =") or paragraph.text.startswith("Расчёт:"):
            _remove_block(paragraph._p)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
