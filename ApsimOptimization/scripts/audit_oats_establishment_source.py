#!/usr/bin/env python3
"""Audit supplied 2023 oat sowing and stand-density evidence."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import tarfile
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OBSERVATIONS = ROOT / "data" / "observations" / "multifield_2019_2023.json"
DEFAULT_MODEL_RESULTS = (
    ROOT / "outputs" / "multifield" / "OATS_WHEAT_RESIDUE"
    / "20260929_120748" / "documented_2022_harvest_export_new_residue"
    / "oats_2023.csv"
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def headers(sheet) -> list[str]:
    return [
        str(cell.value) if cell.value not in (None, "") else f"column_{cell.column}"
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    ]


def rows_for_fields(sheet, year: int, fields: set[int]) -> list[dict]:
    names = headers(sheet)
    index = {name: position for position, name in enumerate(names)}
    if not {"Год", "Поле"} <= set(index):
        raise ValueError(f"{sheet.title}: missing Год/Поле columns")
    result = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = [cell.value for cell in cells]
        if values[index["Год"]] != year or values[index["Поле"]] not in fields:
            continue
        result.append({
            "source_row": row_number,
            **{names[position]: value for position, value in enumerate(values)},
        })
    return result


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row)) or ["empty"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def archive_output(output: Path) -> Path:
    archive_path = ROOT / f"oats_establishment_source_{output.name}.tar.gz"
    files = sorted(file for file in output.iterdir() if file.is_file())
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [{
            "path": file.name,
            "size": file.stat().st_size,
            "sha256": sha256(file),
        } for file in files],
    }
    with tarfile.open(archive_path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.name, recursive=False)
    return archive_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--observations", type=Path, default=DEFAULT_OBSERVATIONS)
    parser.add_argument("--model-results", type=Path, default=DEFAULT_MODEL_RESULTS)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    observations = json.loads(args.observations.read_text(encoding="utf-8"))
    workbook_path = Path(observations["source"]["workbook"])
    actual_hash = sha256(workbook_path)
    if actual_hash != observations["source"].get("workbook_sha256"):
        raise ValueError("source workbook hash does not match extracted observations")

    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output = args.output or (
        ROOT / "outputs" / "multifield" / "OATS_ESTABLISHMENT_SOURCE" / timestamp
    )
    output.mkdir(parents=True, exist_ok=False)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    density_sheet_headers = headers(workbook["Плотность"])
    density_rows = rows_for_fields(workbook["Плотность"], 2023, {2, 7})
    phenology_rows = rows_for_fields(workbook["Фенофазы"], 2023, {2, 7})
    phenology_rows = [
        row for row in phenology_rows
        if "овес" in str(row.get("Культура", "")).casefold().replace("ё", "е")
    ]
    workbook.close()
    write_csv(output / "density_2023_field2_field7.csv", density_rows)
    write_csv(output / "phenology_establishment_2023.csv", phenology_rows)

    model_by_field = {
        row["field"]: row for row in read_csv(args.model_results)
    }
    density_numeric_columns = [
        name for name in (density_rows[0] if density_rows else {})
        if name not in {"source_row", "Год", "Поле", "Точка"}
        and any(
            isinstance(row.get(name), (int, float))
            and math.isfinite(float(row[name]))
            for row in density_rows
        )
    ]
    density_summary = []
    for field in (2, 7):
        field_rows = [row for row in density_rows if row.get("Поле") == field]
        for column in density_numeric_columns:
            values = [
                float(row[column]) for row in field_rows
                if isinstance(row.get(column), (int, float))
                and math.isfinite(float(row[column]))
            ]
            if not values:
                continue
            density_summary.append({
                "field": f"Field{field}",
                "source_column": column,
                "n": len(values),
                "mean": fmean(values),
                "sd": pstdev(values),
                "min": min(values),
                "max": max(values),
                "model_final_plant_population_m2": float(
                    model_by_field[f"Field{field}"]["final_plant_population_m2"]
                ),
                "model_max_stem_population_m2": float(
                    model_by_field[f"Field{field}"]["max_stem_population_m2"]
                ),
            })
    write_csv(output / "density_summary.csv", density_summary)

    phenology_summary = []
    columns = (
        "Норма высева, кг/га", "Всхожесть, %",
        "Плотность посева, млн. зерен/га", "Масса 1000 зёрен, г",
        "Дата посева",
    )
    for row in phenology_rows:
        phenology_summary.append({
            "field": f"Field{int(row['Поле'])}",
            **{column: row.get(column) for column in columns},
        })
    write_csv(output / "phenology_establishment_summary.csv", phenology_summary)

    grain_comparison = []
    for row in phenology_rows:
        field = f"Field{int(row['Поле'])}"
        thousand_grain_mass = row.get("Масса 1000 зёрен, г")
        if not isinstance(thousand_grain_mass, (int, float)):
            continue
        model = model_by_field[field]
        observed_grain_size = float(thousand_grain_mass) / 1000.0
        observed_yield_g_m2 = float(model["observed_yield_t_ha"]) * 100.0
        model_grain_size = float(model["final_grain_size_g"])
        model_grain_number = float(model["final_grain_number_m2"])
        implied_grain_number = observed_yield_g_m2 / observed_grain_size
        grain_comparison.append({
            "field": field,
            "observed_thousand_grain_mass_g": float(thousand_grain_mass),
            "observed_grain_size_g": observed_grain_size,
            "model_grain_size_g": model_grain_size,
            "model_to_observed_grain_size_ratio": (
                model_grain_size / observed_grain_size
            ),
            "model_grain_number_m2": model_grain_number,
            "observed_yield_implied_grain_number_m2": implied_grain_number,
            "implied_to_model_grain_number_ratio": (
                implied_grain_number / model_grain_number
            ),
            "current_grains_per_gram_stem": 45.0,
            "data_derived_grains_per_gram_stem_if_linear": (
                45.0 * implied_grain_number / model_grain_number
            ),
        })
    write_csv(output / "grain_size_and_number_comparison.csv", grain_comparison)

    summary = {
        "experiment": "oats_2023_establishment_source_audit",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "workbook": str(workbook_path),
        "workbook_sha256": actual_hash,
        "density_sheet_headers": density_sheet_headers,
        "density_row_count": len(density_rows),
        "density_summary": density_summary,
        "phenology_establishment": phenology_summary,
        "grain_size_and_number_comparison": grain_comparison,
        "model_results": str(args.model_results),
        "base_model_changed": False,
    }
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    archive = archive_output(output)
    print(json.dumps({
        "output": str(output),
        "archive": str(archive),
        "archive_sha256": sha256(archive),
        "density_summary": density_summary,
        "phenology_establishment": phenology_summary,
        "grain_size_and_number_comparison": grain_comparison,
    }, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
