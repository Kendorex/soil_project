from __future__ import annotations

import argparse
import csv
import math
import os
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


BLUE = "#2E74B5"
OCHRE = "#C69214"
INK = "#243447"
MUTED = "#667085"
GRID = "#D0D5DD"
AXIS = "#98A2B3"


def _font(size: int, bold: bool = False):
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [
        windows_fonts / ("arialbd.ttf" if bold else "arial.ttf"),
        windows_fonts / ("calibrib.ttf" if bold else "calibri.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def _text(draw, xy, value, size=24, fill=INK, bold=False, anchor=None):
    draw.text(xy, str(value), font=_font(size, bold), fill=fill, anchor=anchor)


def _yield_to_t_ha(value: float) -> float:
    # APSIMX Report.Yield in the supplied workbook is in kg/ha.
    return value / 1000.0 if value > 100 else value


def read_simulated_yield(path: Path) -> dict[int, float]:
    sheet = load_workbook(path, read_only=True, data_only=True)["Report"]
    headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    required = {"Clock.Today", "Yield"}
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"Missing Report columns: {', '.join(sorted(missing))}")
    index = {name: position for position, name in enumerate(headers)}
    stage_index = index.get("Wheat.Phenology.CurrentStageName")

    harvest_rows: dict[int, list[float]] = defaultdict(list)
    nonzero_rows: dict[int, list[float]] = defaultdict(list)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        date_value = row[index["Clock.Today"]]
        yield_value = row[index["Yield"]]
        if date_value is None or yield_value is None:
            continue
        numeric_yield = float(yield_value)
        if numeric_yield <= 0:
            continue
        year = date_value.year
        nonzero_rows[year].append(numeric_yield)
        if stage_index is not None and row[stage_index] == "HarvestRipe":
            harvest_rows[year].append(numeric_yield)

    result = {}
    for year in sorted(nonzero_rows):
        candidates = harvest_rows.get(year) or nonzero_rows[year]
        result[year] = _yield_to_t_ha(candidates[-1])
    return result


def read_observed_yield(path: Path) -> dict[int, float]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        rows = csv.DictReader(source)
        return {
            int(row["year"]): float(row["observed_yield_t_ha"])
            for row in rows
            if row.get("year") and row.get("observed_yield_t_ha")
        }


def build_chart(
    simulated: dict[int, float],
    observed: dict[int, float],
    destination: Path,
) -> None:
    years = sorted(set(simulated).intersection(observed))
    if not years:
        raise ValueError("No matching years in simulated and observed data")

    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    plot_left, plot_top, plot_right, plot_bottom = 150, 180, 1535, 735

    _text(
        draw,
        (width / 2, 52),
        "Фактическая и рассчитанная урожайность сорта Zlata",
        36,
        bold=True,
        anchor="ma",
    )
    _text(
        draw,
        (width / 2, 104),
        "Расчёт APSIMX по данным файла «standart 1 field.xlsx»",
        22,
        fill=MUTED,
        anchor="ma",
    )
    _text(draw, (plot_left, plot_top - 38), "Урожайность, т/га", 21, fill=MUTED)

    maximum = max(max(observed[year], simulated[year]) for year in years)
    y_max = max(0.5, math.ceil((maximum + 0.25) * 2) / 2)
    tick_count = int(round(y_max / 0.5))
    for step in range(tick_count + 1):
        value = step * 0.5
        y = plot_bottom - value / y_max * (plot_bottom - plot_top)
        draw.line((plot_left, y, plot_right, y), fill=GRID, width=1)
        _text(draw, (plot_left - 18, y), f"{value:.1f}".replace(".", ","), 18, fill=MUTED, anchor="rm")

    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill=AXIS, width=2)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill=AXIS, width=2)

    slot = (plot_right - plot_left) / len(years)
    bar_width = min(120, slot * 0.26)
    for index, year in enumerate(years):
        center = plot_left + (index + 0.5) * slot
        values = ((observed[year], BLUE, center - bar_width - 8), (simulated[year], OCHRE, center + 8))
        for value, color, x0 in values:
            x1 = x0 + bar_width
            y = plot_bottom - value / y_max * (plot_bottom - plot_top)
            draw.rounded_rectangle((x0, y, x1, plot_bottom), radius=5, fill=color)
            label = f"{value:.2f}".replace(".", ",")
            _text(draw, ((x0 + x1) / 2, y - 14), label, 19, fill=INK, bold=True, anchor="ms")
        _text(draw, (center, plot_bottom + 32), year, 21, fill="#475467", bold=True, anchor="ma")

    legend_y = 820
    legend_items = (("Фактическая", BLUE), ("Рассчитанная APSIMX", OCHRE))
    legend_widths = (230, 330)
    total_width = sum(legend_widths)
    x = (width - total_width) / 2
    for (label, color), item_width in zip(legend_items, legend_widths):
        draw.rounded_rectangle((x, legend_y - 14, x + 28, legend_y + 14), radius=4, fill=color)
        _text(draw, (x + 42, legend_y), label, 20, anchor="lm")
        x += item_width

    destination.parent.mkdir(parents=True, exist_ok=True)
    image.save(destination, format="PNG", optimize=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build observed vs APSIMX yield chart")
    parser.add_argument("source_xlsx", type=Path)
    parser.add_argument("observed_csv", type=Path)
    parser.add_argument("output_png", type=Path)
    args = parser.parse_args()

    simulated = read_simulated_yield(args.source_xlsx)
    observed = read_observed_yield(args.observed_csv)
    build_chart(simulated, observed, args.output_png)
    print("Compared years:", ", ".join(map(str, sorted(set(simulated).intersection(observed)))))
    print("Observed:", observed)
    print("Simulated:", {year: round(value, 6) for year, value in simulated.items()})
    print("Saved:", args.output_png.resolve())


if __name__ == "__main__":
    main()
