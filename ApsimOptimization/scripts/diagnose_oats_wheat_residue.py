#!/usr/bin/env python3
"""Test source-based and extreme 2023 oat surface-residue hypotheses."""
from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import hashlib
import io
import json
import sqlite3
import sys
import tarfile
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.historical_case import load_observations, records_for_year
from src.multifield_model import fingerprint, inventory, prepare, read_json, write_json
from src.multifield_optimizer import ROOT, load_config, run_apsim
from src.oats_field_state import paired_oat_metrics
from src.oats_physiology import collect_oats_diagnostics, patch_oats_diagnostic_reports
from src.oats_preseason_n import collect_preseason_rows, build_preseason_n_summary
from src.oats_residue_sensitivity import (
    add_full_residue_removal, add_pre_sowing_residue_cap,
    defer_wheat_auto_harvest_for_field_years,
    patch_export_new_residue_at_cereal_to_oats_harvest,
)


FIELDS = ("Field2", "Field7")
CANDIDATES = (
    "current_residue_retention",
    "documented_2022_wheat_harvest_timing",
    "documented_2022_harvest_export_new_residue",
    "source_2023_stubble_mean_as_pre_sowing_cap",
    "remove_all_surface_residue_after_2022_wheat",
)


def source_residue_targets(observations: dict) -> tuple[dict[str, float], list[dict]]:
    """Read field means from the supplied 2023 stubble-residue observations."""
    source = observations["source"]
    workbook_path = Path(source["workbook"])
    actual_hash = fingerprint(workbook_path)
    if actual_hash != source.get("workbook_sha256"):
        raise ValueError("source workbook hash does not match extracted observations")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Растительные_остатки"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value) for value in next(rows)]
    index = {name: position for position, name in enumerate(header)}
    raw_name = "Растительные остатки (пожнивные), ц/га"
    corrected_name = "Растительные остатки испр./скорр. (пожнивные), ц/га"
    required = {"Год", "Поле", "Точка", raw_name, corrected_name}
    if not required <= set(index):
        raise ValueError(f"residue source sheet misses columns: {sorted(required - set(index))}")
    values: dict[str, list[tuple[float, str]]] = {field: [] for field in FIELDS}
    for row in rows:
        if row[index["Год"]] != 2023:
            continue
        field = f"Field{int(row[index['Поле']])}"
        if field not in values:
            continue
        corrected = row[index[corrected_name]]
        raw = row[index[raw_name]]
        selected = corrected if isinstance(corrected, (int, float)) else raw
        if not isinstance(selected, (int, float)):
            continue
        source_column = corrected_name if isinstance(corrected, (int, float)) else raw_name
        values[field].append((float(selected), source_column))
    workbook.close()

    targets: dict[str, float] = {}
    summaries = []
    for field in FIELDS:
        field_values = [value for value, _ in values[field]]
        if not field_values:
            raise ValueError(f"{field}: no numeric 2023 stubble-residue observations")
        mean_c_ha = fmean(field_values)
        targets[field] = mean_c_ha * 100.0
        summaries.append({
            "field": field,
            "year": 2023,
            "n_points": len(field_values),
            "mean_stubble_c_ha": mean_c_ha,
            "sd_stubble_c_ha": pstdev(field_values),
            "min_stubble_c_ha": min(field_values),
            "max_stubble_c_ha": max(field_values),
            "target_surface_residue_kg_ha": targets[field],
            "source_columns": ";".join(sorted({column for _, column in values[field]})),
            "timing_status": "sampling date not stated in source sheet",
        })
    return targets, summaries


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def removal_events(database: Path, meta: dict, candidate: str) -> list[dict]:
    output = []
    uri = database.resolve().as_uri() + "?mode=ro"
    with sqlite3.connect(uri, uri=True) as connection:
        connection.row_factory = sqlite3.Row
        simulation_id = connection.execute(
            'SELECT ID FROM "_Simulations" WHERE Name=?', (meta["simulation"],)
        ).fetchone()[0]
        for field in FIELDS:
            table = f"Report{field.removeprefix('Field')}"
            columns = {row[1] for row in connection.execute(f'PRAGMA table_info("{table}")')}
            names = {"ResidueRemovalTriggered", "ResidueRemovedMassKgHa", "ResidueRemovedNKgHa"}
            cereal_names = {
                "CerealResidueRemovalCount", "CerealResidueRemovedMassKgHa",
                "CerealResidueRemovedNKgHa",
            }
            if cereal_names <= columns:
                row = connection.execute(
                    f'''SELECT max(CerealResidueRemovalCount) AS Triggered,
                        max(CerealResidueRemovedMassKgHa) AS RemovedMass,
                        max(CerealResidueRemovedNKgHa) AS RemovedN
                        FROM "{table}" WHERE SimulationID=?''', (simulation_id,)
                ).fetchone()
                output.append({
                    "candidate": candidate, "field": field,
                    "triggered": float(row["Triggered"] or 0),
                    "removed_mass_kg_ha": float(row["RemovedMass"] or 0),
                    "removed_n_kg_ha": float(row["RemovedN"] or 0),
                    "status": "documented_harvest_new_residue_export_hypothesis",
                })
                continue
            if not names <= columns:
                output.append({
                    "candidate": candidate, "field": field,
                    "triggered": 0, "removed_mass_kg_ha": 0, "removed_n_kg_ha": 0,
                    "status": "control_without_removal_manager",
                })
                continue
            optional = {
                "ResidueTargetMassKgHa": "TargetMass",
                "ResidueBeforeMassKgHa": "BeforeMass",
                "ResidueAfterMassKgHa": "AfterMass",
            }
            optional_sql = ", ".join(
                f'max("{column}") AS "{alias}"'
                for column, alias in optional.items() if column in columns
            )
            if optional_sql:
                optional_sql = ", " + optional_sql
            row = connection.execute(
                f'''SELECT max(ResidueRemovalTriggered) AS Triggered,
                    max(ResidueRemovedMassKgHa) AS RemovedMass,
                    max(ResidueRemovedNKgHa) AS RemovedN{optional_sql}
                    FROM "{table}" WHERE SimulationID=?''', (simulation_id,)
            ).fetchone()
            event = {
                "candidate": candidate, "field": field,
                "triggered": float(row["Triggered"] or 0),
                "removed_mass_kg_ha": float(row["RemovedMass"] or 0),
                "removed_n_kg_ha": float(row["RemovedN"] or 0),
                "status": (
                    "source_mean_pre_sowing_cap_hypothesis"
                    if candidate == "source_2023_stubble_mean_as_pre_sowing_cap"
                    else "diagnostic_complete_removal"
                ),
            }
            for column, alias in optional.items():
                if column in columns:
                    key = {
                        "TargetMass": "target_mass_kg_ha",
                        "BeforeMass": "before_mass_kg_ha",
                        "AfterMass": "after_mass_kg_ha",
                    }[alias]
                    event[key] = float(row[alias] or 0)
            output.append(event)
    return output


def create_archive(path: Path, root: Path) -> None:
    allowed = {
        "summary.json", "candidate_metrics.csv", "case_predictions.csv",
        "residue_events.csv", "source_residue_targets.csv", "candidate_summary.json", "daily_preseason_n.csv",
        "phase_snapshots.csv", "interval_balance.csv", "oats_2023.csv",
        "report_schema.json", "config_snapshot.json", "model_inventory.json",
        "preparation.json", "apsim.log", "changes.json",
    }
    files = sorted(file for file in root.rglob("*")
                   if file.is_file() and file.name in allowed)
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [{
            "path": file.relative_to(root).as_posix(), "size": file.stat().st_size,
            "sha256": hashlib.sha256(file.read_bytes()).hexdigest(),
        } for file in files],
    }
    with tarfile.open(path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.relative_to(root).as_posix(), recursive=False)


def run_candidate(config: dict, observations: dict, windows: dict,
                  candidate: str, output: Path,
                  residue_targets: dict[str, float]) -> tuple[dict, list[dict], list[dict]]:
    output.mkdir(parents=True, exist_ok=False)
    data = read_json(config["model"])
    report_changes = patch_oats_diagnostic_reports(data, FIELDS)
    changes = []
    if candidate == "remove_all_surface_residue_after_2022_wheat":
        changes = add_full_residue_removal(data, "Wheat", FIELDS)
    elif candidate == "documented_2022_wheat_harvest_timing":
        changes = defer_wheat_auto_harvest_for_field_years(
            data, {field: {2022} for field in FIELDS}
        )
    elif candidate == "documented_2022_harvest_export_new_residue":
        changes = defer_wheat_auto_harvest_for_field_years(
            data, {field: {2022} for field in FIELDS}
        )
        changes.extend({"rotation_manager_path": path} for path in
                       patch_export_new_residue_at_cereal_to_oats_harvest(
                           data, FIELDS, {field: {2022} for field in FIELDS}
                       ))
    elif candidate == "source_2023_stubble_mean_as_pre_sowing_cap":
        changes = add_pre_sowing_residue_cap(
            data, "Wheat", residue_targets, date="2023-05-03"
        )
    source = output / "source_candidate.apsimx"
    write_json(source, data)
    case = copy.deepcopy(config)
    case["model"] = str(source.resolve())
    case["model_provenance"] = None
    case["scenario"] = {
        "kind": "diagnostic", "name": candidate, "recommendation_status": "prohibited",
    }
    meta = inventory(case)
    write_json(output / "config_snapshot.json", case)
    write_json(output / "model_inventory.json", meta)
    prepared = prepare(case, output)
    database = run_apsim(case, prepared)
    daily = collect_preseason_rows(database, meta, windows)
    snapshots, intervals = build_preseason_n_summary(daily, windows)
    _, oats, schema = collect_oats_diagnostics(
        database, meta, records_for_year(observations, 2023), year=2023,
        scenario=candidate,
    )
    metrics = paired_oat_metrics(oats, 2023, *FIELDS)
    events = removal_events(database, meta, candidate)
    write_csv(output / "daily_preseason_n.csv", daily)
    write_csv(output / "phase_snapshots.csv", snapshots)
    write_csv(output / "interval_balance.csv", intervals)
    write_csv(output / "oats_2023.csv", oats)
    write_csv(output / "residue_events.csv", events)
    write_json(output / "report_schema.json", schema)
    write_json(output / "changes.json", {
        "candidate": candidate, "changes": changes,
        "diagnostic_report_changes": report_changes, "base_model_changed": False,
    })
    result = {"candidate": candidate, "yield_metrics": metrics, "events": events}
    write_json(output / "candidate_summary.json", result)
    return result, oats, intervals


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", action="store_true")
    parser.add_argument("--events-database", type=Path)
    parser.add_argument("--events-candidate", default="database_event_audit")
    parser.add_argument(
        "--config", type=Path,
        default=ROOT / "configs" / "multifield_historical_replay.yaml",
    )
    args = parser.parse_args(argv)
    config = load_config(args.config)
    if args.events_database:
        meta = inventory(config)
        print(json.dumps(
            removal_events(args.events_database, meta, args.events_candidate),
            ensure_ascii=False, indent=2,
        ))
        return 0
    if not args.run:
        parser.error("one of --run or --events-database is required")
    observations = load_observations(config["historical_observations"])
    residue_targets, residue_target_rows = source_residue_targets(observations)
    previous = {row["field"]: row for row in records_for_year(observations, 2022)}
    current = {row["field"]: row for row in records_for_year(observations, 2023)}
    windows = {}
    for field in FIELDS:
        harvest = previous[field].get("observed_harvest_date")
        sowing = current[field].get("observed_sow_date")
        if harvest != "2022-08-17" or not sowing:
            raise ValueError(f"{field}: expected documented 2022 wheat harvest date")
        windows[field] = {
            "previous_harvest_date": harvest, "sowing_date": sowing,
            "end_date": "2023-12-31",
        }

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    root = ROOT / "outputs" / "multifield" / "OATS_WHEAT_RESIDUE" / stamp
    root.mkdir(parents=True, exist_ok=False)
    write_csv(root / "source_residue_targets.csv", residue_target_rows)
    results, predictions, interval_rows, event_rows = [], [], [], []
    for index, candidate in enumerate(CANDIDATES, 1):
        print(f"[{index}/{len(CANDIDATES)}] {candidate}", flush=True)
        result, oats, intervals = run_candidate(
            config, observations, windows, candidate, root / candidate,
            residue_targets,
        )
        results.append(result)
        predictions.extend({"candidate": candidate, **row} for row in oats)
        interval_rows.extend({"candidate": candidate, **row} for row in intervals)
        event_rows.extend(result["events"])
        metric = result["yield_metrics"]
        print(
            f"  Field2={metric['fields_detail']['Field2']['simulated_yield_t_ha']:.3f}; "
            f"Field7={metric['fields_detail']['Field7']['simulated_yield_t_ha']:.3f}; "
            f"RMSE={metric['rmse_t_ha']:.3f}", flush=True,
        )

    control, documented_timing, documented_export, source_cap, removal = results
    metric_rows = []
    for result in results:
        metric = result["yield_metrics"]
        metric_rows.append({
            "candidate": result["candidate"], "mae_t_ha": metric["mae_t_ha"],
            "rmse_t_ha": metric["rmse_t_ha"], "bias_t_ha": metric["bias_t_ha"],
            "field2_yield_t_ha": metric["fields_detail"]["Field2"]["simulated_yield_t_ha"],
            "field7_yield_t_ha": metric["fields_detail"]["Field7"]["simulated_yield_t_ha"],
            "ordering_matches_observations": metric["ordering_matches_observations"],
        })
    write_csv(root / "candidate_metrics.csv", metric_rows)
    write_csv(root / "case_predictions.csv", predictions)
    write_csv(root / "interval_balance.csv", interval_rows)
    write_csv(root / "residue_events.csv", event_rows)
    write_json(root / "summary.json", {
        "experiment": "2023_oats_source_based_and_extreme_surface_residue_hypotheses",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": "engineering_diagnostic_only", "base_model_changed": False,
        "candidates": results,
        "rmse_change_t_ha": (
            removal["yield_metrics"]["rmse_t_ha"] - control["yield_metrics"]["rmse_t_ha"]
        ),
        "source_cap_rmse_change_t_ha": (
            source_cap["yield_metrics"]["rmse_t_ha"] - control["yield_metrics"]["rmse_t_ha"]
        ),
        "documented_harvest_timing_rmse_change_t_ha": (
            documented_timing["yield_metrics"]["rmse_t_ha"]
            - control["yield_metrics"]["rmse_t_ha"]
        ),
        "documented_harvest_export_rmse_change_t_ha": (
            documented_export["yield_metrics"]["rmse_t_ha"]
            - control["yield_metrics"]["rmse_t_ha"]
        ),
        "inputs": {
            "model": config["model"], "model_sha256": fingerprint(config["model"]),
            "weather": config["weather"], "weather_sha256": fingerprint(config["weather"]),
            "observations": config["historical_observations"],
            "observations_sha256": fingerprint(config["historical_observations"]),
            "source_workbook": observations["source"]["workbook"],
            "source_workbook_sha256": observations["source"]["workbook_sha256"],
        },
        "interpretation_limits": [
            "Complete surface-residue removal is an extreme mechanism test, not a management recommendation.",
            "The source workbook contains 2023 stubble-residue mass but does not state its sampling date.",
            "Using the 2023 field mean as a pre-sowing cap is an explicit timing hypothesis, not a confirmed historical operation.",
            "The source data do not state whether 2022 wheat straw was retained, baled, burned, or incorporated.",
            "The documented-timing candidate suppresses readiness-based Wheat harvest only in 2022 on Field2 and Field7.",
            "The diagnostic removes all surface organic matter present after the documented 2022 wheat harvest.",
            "No fertilizer, soil, crop, nutrient-target, or optimization bound is changed.",
        ],
    })
    archive = ROOT / f"oats_wheat_residue_{stamp}.tar.gz"
    create_archive(archive, root)
    print(f"OUTPUT: {root}", flush=True)
    print(f"SEND THIS ARCHIVE: {archive}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
