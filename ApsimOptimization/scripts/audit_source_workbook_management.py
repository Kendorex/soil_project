#!/usr/bin/env python3
"""Audit the source workbook for residue and field-management evidence.

The script is read-only with respect to the workbook.  It records every cell
whose text contains one of the configured Russian management terms, together
with the non-empty values from the same row.  The result is intended to show
whether a historical APSIM assumption is supported by the supplied source.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import tarfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OBSERVATIONS = ROOT / "data" / "observations" / "multifield_2019_2023.json"

TERM_GROUPS = {
    "residue": ("солом", "остат", "стерн", "пожнив", "мульч"),
    "residue_action": ("задел", "вывоз", "измельч"),
    "tillage": ("обработк", "вспаш", "диск", "плуг"),
    "generic_management": ("уборк", "севооборот", "удобрен", "технолог"),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def source_workbook(observations_path: Path) -> Path:
    payload = json.loads(observations_path.read_text(encoding="utf-8"))
    return Path(payload["source"]["workbook"])


def terms_in(value: object) -> list[tuple[str, str]]:
    text = str(value).casefold()
    return [
        (group, term)
        for group, terms in TERM_GROUPS.items()
        for term in terms
        if term in text
    ]


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or ["empty"])
        writer.writeheader()
        writer.writerows(rows)


def create_archive(path: Path, output: Path) -> None:
    files = sorted(file for file in output.iterdir() if file.is_file())
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [
            {
                "path": file.name,
                "size": file.stat().st_size,
                "sha256": sha256(file),
            }
            for file in files
        ],
    }
    with tarfile.open(path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.name, recursive=False)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--observations", type=Path, default=DEFAULT_OBSERVATIONS)
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args(argv)

    workbook_path = args.workbook or source_workbook(args.observations)
    if not workbook_path.is_file():
        parser.error(f"workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    structure: list[dict] = []
    hits: list[dict] = []
    residue_rows: list[dict] = []
    group_counts = {group: 0 for group in TERM_GROUPS}

    for sheet in workbook.worksheets:
        nonempty = 0
        text_cells = 0
        header: list[str] | None = None
        for row in sheet.iter_rows():
            row_values = {
                cell.coordinate: cell.value
                for cell in row
                if cell.value not in (None, "")
            }
            nonempty += len(row_values)
            text_cells += sum(isinstance(value, str) for value in row_values.values())
            if sheet.title == "Растительные_остатки":
                values = [cell.value for cell in row]
                if header is None:
                    header = [str(value) if value not in (None, "") else f"column_{index}"
                              for index, value in enumerate(values, start=1)]
                elif any(value not in (None, "") for value in values):
                    residue_rows.append({
                        "source_row": row[0].row,
                        **{name: value for name, value in zip(header, values)},
                    })
            row_context = json.dumps(row_values, ensure_ascii=False, default=str)
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                matches = terms_in(cell.value)
                if not matches:
                    continue
                groups = sorted({group for group, _ in matches})
                matched_terms = sorted({term for _, term in matches})
                for group in groups:
                    group_counts[group] += 1
                hits.append({
                    "sheet": sheet.title,
                    "cell": cell.coordinate,
                    "value": cell.value,
                    "groups": ";".join(groups),
                    "matched_terms": ";".join(matched_terms),
                    "row_context_json": row_context,
                })
        structure.append({
            "sheet": sheet.title,
            "state": sheet.sheet_state,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "nonempty_cells": nonempty,
            "text_cells": text_cells,
        })
    workbook.close()

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    output = ROOT / "outputs" / "multifield" / "SOURCE_MANAGEMENT_AUDIT" / stamp
    output.mkdir(parents=True, exist_ok=False)
    write_csv(output / "workbook_structure.csv", structure)
    write_csv(output / "management_term_hits.csv", hits)
    write_csv(output / "source_residue_rows.csv", residue_rows)

    residue_specific_hits = sum(
        1 for row in hits
        if "residue" in row["groups"] or "residue_action" in row["groups"]
    )
    summary = {
        "experiment": "source_workbook_management_evidence_audit",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": "input_evidence_audit",
        "base_model_changed": False,
        "workbook": str(workbook_path),
        "workbook_sha256": sha256(workbook_path),
        "sheet_count": len(structure),
        "term_groups": TERM_GROUPS,
        "hit_count": len(hits),
        "group_hit_counts": group_counts,
        "residue_specific_hit_count": residue_specific_hits,
        "source_residue_data_rows": len(residue_rows),
        "interpretation": (
            "Residue-specific text was found; inspect management_term_hits.csv."
            if residue_specific_hits
            else "No residue-specific text was found in workbook cell values."
        ),
        "limits": [
            "The audit searches cell values, not embedded images or external linked files.",
            "Generic harvest and technology labels do not establish residue handling.",
            "Absence of a term means the workbook does not document the practice; it does not prove the practice did not occur.",
        ],
    }
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    archive = ROOT / f"source_management_audit_{stamp}.tar.gz"
    create_archive(archive, output)

    print(f"WORKBOOK SHA256: {summary['workbook_sha256']}")
    print(f"SHEETS: {summary['sheet_count']}")
    print(f"ALL TERM HITS: {summary['hit_count']}")
    print(f"RESIDUE-SPECIFIC HITS: {summary['residue_specific_hit_count']}")
    print(f"OUTPUT: {output}")
    print(f"ARCHIVE: {archive}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
