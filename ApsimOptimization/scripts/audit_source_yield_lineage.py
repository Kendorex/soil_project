#!/usr/bin/env python3
"""Trace 2023 oat yield values back to cells in the supplied workbook."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import tarfile
from pathlib import Path
from statistics import fmean, pstdev

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OBSERVATIONS = ROOT / "data" / "observations" / "multifield_2019_2023.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row)) or ["empty"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def archive_output(output: Path) -> Path:
    archive_path = ROOT / f"source_yield_lineage_{output.name}.tar.gz"
    files = sorted(file for file in output.iterdir() if file.is_file())
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [{
            "path": file.name,
            "size": file.stat().st_size,
            "sha256": sha256(file),
        } for file in files],
    }
    with tarfile.open(archive_path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.name, recursive=False)
    return archive_path


def normalized(value) -> str:
    return str(value or "").strip().casefold().replace("ё", "е")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--observations", type=Path, default=DEFAULT_OBSERVATIONS)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    observations = json.loads(args.observations.read_text(encoding="utf-8"))
    workbook_path = Path(observations["source"]["workbook"])
    actual_hash = sha256(workbook_path)
    expected_hash = observations["source"].get("workbook_sha256")
    if actual_hash != expected_hash:
        raise ValueError("source workbook hash does not match extracted observations")

    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output = args.output or (
        ROOT / "outputs" / "multifield" / "SOURCE_YIELD_LINEAGE" / timestamp
    )
    output.mkdir(parents=True, exist_ok=False)

    formulas = load_workbook(workbook_path, read_only=True, data_only=False)
    values = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet_formula = formulas["Фенофазы"]
    sheet_values = values["Фенофазы"]
    header = [cell.value for cell in next(sheet_values.iter_rows(min_row=1, max_row=1))]
    index = {str(name): position + 1 for position, name in enumerate(header)}
    required = {"Год", "Поле", "Культура", "Сорт", "Средняя урожайность, т/га"}
    if not required <= set(index):
        raise ValueError(f"Фенофазы misses columns: {sorted(required - set(index))}")

    selected_rows = []
    target_values = set()
    for row_number in range(2, sheet_values.max_row + 1):
        year = sheet_values.cell(row_number, index["Год"]).value
        field = sheet_values.cell(row_number, index["Поле"]).value
        crop = sheet_values.cell(row_number, index["Культура"]).value
        if year != 2023 or field not in (2, 7) or "овес" not in normalized(crop):
            continue
        yield_column = index["Средняя урожайность, т/га"]
        value_cell = sheet_values.cell(row_number, yield_column)
        formula_cell = sheet_formula.cell(row_number, yield_column)
        value = value_cell.value
        if isinstance(value, (int, float)) and math.isfinite(float(value)):
            target_values.add(float(value))
        selected_rows.append({
            "sheet": "Фенофазы",
            "row": row_number,
            "field": f"Field{int(field)}",
            "year": int(year),
            "crop": crop,
            "cultivar": sheet_values.cell(row_number, index["Сорт"]).value,
            "yield_cell": value_cell.coordinate,
            "yield_cached_value_t_ha": value,
            "yield_formula_or_literal": formula_cell.value,
            "cell_data_type": formula_cell.data_type,
            "number_format": formula_cell.number_format,
        })
    if len(selected_rows) != 2:
        raise ValueError(f"expected two 2023 oat yield rows, found {len(selected_rows)}")
    write_csv(output / "selected_2023_oat_yield_cells.csv", selected_rows)

    occurrences = []
    keyword_cells = []
    for sheet_name in values.sheetnames:
        value_sheet = values[sheet_name]
        formula_sheet = formulas[sheet_name]
        for row in value_sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, (int, float)) and any(
                    math.isclose(float(value), target, rel_tol=0, abs_tol=1e-9)
                    for target in target_values
                ):
                    formula_cell = formula_sheet[cell.coordinate]
                    occurrences.append({
                        "sheet": sheet_name,
                        "cell": cell.coordinate,
                        "cached_value": value,
                        "formula_or_literal": formula_cell.value,
                        "cell_data_type": formula_cell.data_type,
                    })
                if isinstance(value, str) and "урожайн" in normalized(value):
                    keyword_cells.append({
                        "sheet": sheet_name,
                        "cell": cell.coordinate,
                        "text": value,
                    })
    write_csv(output / "matching_numeric_cells.csv", occurrences)
    write_csv(output / "yield_keyword_cells.csv", keyword_cells)

    raw_yield_sheet = values["Урожайность"]
    raw_headers = [
        str(cell.value) if cell.value not in (None, "") else f"column_{cell.column}"
        for cell in next(raw_yield_sheet.iter_rows(min_row=1, max_row=1))
    ]
    raw_rows = []
    for row_number, cells in enumerate(
        raw_yield_sheet.iter_rows(min_row=2), start=2
    ):
        row_values = [cell.value for cell in cells]
        if 2023 not in row_values:
            continue
        raw_rows.append({
            "source_row": row_number,
            **{
                raw_headers[position]: value
                for position, value in enumerate(row_values)
            },
        })
    write_csv(output / "raw_yield_sheet_2023_rows.csv", raw_rows)
    raw_comparison = []
    for selected in selected_rows:
        field_number = int(selected["field"].removeprefix("Field"))
        raw_values = [
            float(row["Урожайность, ц/га"])
            for row in raw_rows
            if row.get("Поле") == field_number
            and isinstance(row.get("Урожайность, ц/га"), (int, float))
        ]
        if not raw_values:
            raise ValueError(f"no raw 2023 yield samples for {selected['field']}")
        raw_mean_t_ha = fmean(raw_values) / 10.0
        summary_value = float(selected["yield_cached_value_t_ha"])
        raw_comparison.append({
            "field": selected["field"],
            "summary_yield_t_ha": summary_value,
            "raw_n": len(raw_values),
            "raw_mean_t_ha": raw_mean_t_ha,
            "raw_sd_t_ha": pstdev(raw_values) / 10.0,
            "raw_min_t_ha": min(raw_values) / 10.0,
            "raw_max_t_ha": max(raw_values) / 10.0,
            "summary_minus_raw_mean_t_ha": summary_value - raw_mean_t_ha,
            "summary_minus_raw_mean_percent": (
                100.0 * (summary_value - raw_mean_t_ha) / raw_mean_t_ha
            ),
        })
    write_csv(output / "yield_summary_vs_raw_samples.csv", raw_comparison)

    summary = {
        "experiment": "source_yield_lineage_audit",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "workbook": str(workbook_path),
        "workbook_sha256": actual_hash,
        "selected_rows": selected_rows,
        "matching_numeric_cell_count": len(occurrences),
        "matching_numeric_cells": occurrences,
        "yield_keyword_cell_count": len(keyword_cells),
        "raw_yield_sheet_headers": raw_headers,
        "raw_yield_sheet_2023_row_count": len(raw_rows),
        "yield_summary_vs_raw_samples": raw_comparison,
        "formula_status": {
            row["field"]: (
                "formula" if str(row["yield_formula_or_literal"]).startswith("=")
                else "hard_coded_numeric_value"
            )
            for row in selected_rows
        },
        "explicit_moisture_basis_found_in_yield_keyword_cells": any(
            "влаж" in normalized(row["text"]) for row in keyword_cells
        ),
        "interpretation_limit": (
            "Absence of a moisture note in workbook labels does not prove that the "
            "source measurements were not moisture-corrected."
        ),
    }
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    formulas.close()
    values.close()
    archive = archive_output(output)
    print(json.dumps({
        "output": str(output),
        "archive": str(archive),
        "archive_sha256": sha256(archive),
        "selected_rows": selected_rows,
        "matching_numeric_cell_count": len(occurrences),
        "yield_keyword_cell_count": len(keyword_cells),
        "raw_yield_sheet_headers": raw_headers,
        "raw_yield_sheet_2023_row_count": len(raw_rows),
        "yield_summary_vs_raw_samples": raw_comparison,
        "explicit_moisture_basis_found_in_yield_keyword_cells": summary[
            "explicit_moisture_basis_found_in_yield_keyword_cells"
        ],
    }, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
