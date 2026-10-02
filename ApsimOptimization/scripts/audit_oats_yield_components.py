#!/usr/bin/env python3
"""Compare observed and simulated oat grain size across historical cases."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import tarfile
from pathlib import Path
from statistics import fmean

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OBSERVATIONS = ROOT / "data" / "observations" / "multifield_2019_2023.json"
DEFAULT_HISTORY = (
    ROOT / "outputs" / "multifield" / "OATS_SOURCE_FACTORS"
    / "20260928_112031" / "current_yakov" / "oats_history.csv"
)
DEFAULT_TRANSITION = (
    ROOT / "outputs" / "multifield" / "OATS_WHEAT_RESIDUE"
    / "20260929_120748" / "documented_2022_harvest_export_new_residue"
    / "oats_2023.csv"
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row)) or ["empty"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def source_grain_sizes(observations: dict) -> dict[tuple[int, str], dict]:
    workbook_path = Path(observations["source"]["workbook"])
    if sha256(workbook_path) != observations["source"].get("workbook_sha256"):
        raise ValueError("source workbook hash does not match observations")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Фенофазы"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {str(name): position for position, name in enumerate(header)}
    result = {}
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        crop = str(cells[index["Культура"]] or "").casefold().replace("ё", "е")
        if "овес" not in crop:
            continue
        year = cells[index["Год"]]
        field = cells[index["Поле"]]
        mass = cells[index["Масса 1000 зёрен, г"]]
        if not isinstance(year, int) or not isinstance(field, int):
            continue
        result[(year, f"Field{field}")] = {
            "observed_thousand_grain_mass_g": (
                float(mass) if isinstance(mass, (int, float)) else None
            ),
            "source_cultivar": cells[index["Сорт"]],
        }
    workbook.close()
    return result


def comparison_row(row: dict[str, str], scenario: str, source: dict) -> dict | None:
    observed_tkw = source.get("observed_thousand_grain_mass_g")
    if observed_tkw is None:
        return None
    observed_size = observed_tkw / 1000.0
    simulated_size = float(row["final_grain_size_g"])
    observed_yield = float(row["observed_yield_t_ha"])
    simulated_number = float(row["final_grain_number_m2"])
    implied_number = observed_yield * 100.0 / observed_size
    return {
        "scenario": scenario,
        "year": int(row["year"]),
        "field": row["field"],
        "source_cultivar": source["source_cultivar"],
        "observed_yield_t_ha": observed_yield,
        "simulated_yield_t_ha": float(row["simulated_yield_t_ha"]),
        "observed_thousand_grain_mass_g": observed_tkw,
        "simulated_thousand_grain_mass_g": simulated_size * 1000.0,
        "simulated_to_observed_grain_mass_ratio": simulated_size / observed_size,
        "observed_yield_implied_grain_number_m2": implied_number,
        "simulated_grain_number_m2": simulated_number,
        "simulated_to_implied_grain_number_ratio": simulated_number / implied_number,
        "grain_n_allocation_to_demand_ratio": float(
            row["grain_n_allocation_to_demand_ratio"]
        ),
        "simulated_harvest_index": float(row["simulated_harvest_index"]),
    }


def group_metrics(rows: list[dict]) -> dict:
    return {
        "n": len(rows),
        "mean_observed_thousand_grain_mass_g": fmean(
            row["observed_thousand_grain_mass_g"] for row in rows
        ) if rows else None,
        "mean_simulated_thousand_grain_mass_g": fmean(
            row["simulated_thousand_grain_mass_g"] for row in rows
        ) if rows else None,
        "mean_simulated_to_observed_grain_mass_ratio": fmean(
            row["simulated_to_observed_grain_mass_ratio"] for row in rows
        ) if rows else None,
        "mean_simulated_to_implied_grain_number_ratio": fmean(
            row["simulated_to_implied_grain_number_ratio"] for row in rows
        ) if rows else None,
    }


def archive_output(output: Path) -> Path:
    archive_path = ROOT / f"oats_yield_components_{output.name}.tar.gz"
    files = sorted(file for file in output.iterdir() if file.is_file())
    manifest = {
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "files": [{
            "path": file.name, "size": file.stat().st_size,
            "sha256": sha256(file),
        } for file in files],
    }
    with tarfile.open(archive_path, "w:gz") as archive:
        payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        info = tarfile.TarInfo("manifest.json")
        info.size = len(payload)
        info.mtime = int(dt.datetime.now().timestamp())
        archive.addfile(info, io.BytesIO(payload))
        for file in files:
            archive.add(file, arcname=file.name, recursive=False)
    return archive_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--observations", type=Path, default=DEFAULT_OBSERVATIONS)
    parser.add_argument("--history", type=Path, default=DEFAULT_HISTORY)
    parser.add_argument("--transition", type=Path, default=DEFAULT_TRANSITION)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    observations = json.loads(args.observations.read_text(encoding="utf-8"))
    source = source_grain_sizes(observations)
    baseline = []
    for row in read_csv(args.history):
        key = (int(row["year"]), row["field"])
        if key not in source or float(row["simulated_yield_t_ha"]) <= 0:
            continue
        item = comparison_row(row, "current_historical_model", source[key])
        if item:
            baseline.append(item)
    transition = []
    for row in read_csv(args.transition):
        key = (int(row["year"]), row["field"])
        item = comparison_row(row, "2023_transition_diagnostic", source[key])
        if item:
            transition.append(item)

    rows = baseline + transition
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output = args.output or (
        ROOT / "outputs" / "multifield" / "OATS_YIELD_COMPONENTS" / timestamp
    )
    output.mkdir(parents=True, exist_ok=False)
    write_csv(output / "yield_component_comparison.csv", rows)
    groups = {
        "yakov_development_2020_2022": group_metrics([
            row for row in baseline
            if 2020 <= row["year"] <= 2022 and row["source_cultivar"] == "Яков"
        ]),
        "yakov_development_excluding_known_field1_2021_lifecycle_failure": group_metrics([
            row for row in baseline
            if 2020 <= row["year"] <= 2022 and row["source_cultivar"] == "Яков"
            and (row["year"], row["field"]) != (2021, "Field1")
        ]),
        "historical_yield_near_matches_absolute_error_le_0_2": group_metrics([
            row for row in baseline
            if 2020 <= row["year"] <= 2022 and row["source_cultivar"] == "Яков"
            and abs(row["simulated_yield_t_ha"] - row["observed_yield_t_ha"]) <= 0.2
        ]),
        "yakov_validation_2023_baseline": group_metrics([
            row for row in baseline
            if row["year"] == 2023 and row["source_cultivar"] == "Яков"
        ]),
        "yakov_validation_2023_transition": group_metrics(transition),
    }
    summary = {
        "experiment": "historical_oat_yield_component_audit",
        "created_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "groups": groups, "base_model_changed": False,
        "interpretation_limits": [
            "Observed yield moisture basis is not stated.",
            "Observed grain number is inferred from observed yield and thousand-grain mass; it is not independently measured.",
            "The transition scenario assumes complete export of fresh 2022 wheat harvest residue.",
        ],
    }
    (output / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    archive = archive_output(output)
    print(json.dumps({
        "output": str(output), "archive": str(archive),
        "archive_sha256": sha256(archive), "groups": groups,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
