from __future__ import annotations

import csv
import copy
import hashlib
import itertools
import json
import math
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from .apsim_editor import FixedParameters, _named, _parameters, _walk, create_working_copy
from .apsim_runner import find_db_for_apsimx, run_apsim_files
from .config import ConfigError, _parse_simple_yaml
from .results_reader import convert_yield_to_t_ha

try:
    import yaml
except ImportError:
    yaml = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class CalibrationError(RuntimeError):
    """Raised when cultivar calibration inputs or APSIMX results are invalid."""


@dataclass(frozen=True)
class CalibrationObservation:
    year: int
    field: int
    cultivar: str
    technology: str
    observed_yield_t_ha: float
    sowing_date: date
    heading_start_date: date
    heading_end_date: date | None
    harvest_date: date | None
    thousand_seed_weight_g: float | None
    sowing_rate_kg_ha: float | None
    germination_percent: float | None
    seed_density_million_ha: float | None
    population_plants_m2: float
    dataset_role: str


def load_calibration_config(path: Path, root: Path) -> dict[str, Any]:
    if not path.is_file():
        raise ConfigError(f"Cultivar calibration config not found: {path}")
    text = path.read_text(encoding="utf-8")
    loaded = yaml.safe_load(text) if yaml is not None else _parse_simple_yaml(text)
    if not isinstance(loaded, dict) or not isinstance(loaded.get("calibration"), dict):
        raise ConfigError("Calibration YAML must contain a calibration mapping")
    loaded["_root"] = root
    loaded["_config_path"] = path
    return loaded


def _number(value: Any) -> float | None:
    if value in (None, "", "Нет данных"):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _date_from_excel(value: Any, year: int) -> date | None:
    if value in (None, "", "Нет данных"):
        return None
    if isinstance(value, datetime):
        return value.date().replace(year=year)
    if isinstance(value, date):
        return value.replace(year=year)
    text = str(value).strip().replace(",", ".")
    for separator in (".", "/", "-"):
        parts = text.split(separator)
        if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
            return date(year, int(parts[1]), int(parts[0]))
    raise CalibrationError(f"Cannot parse field date {value!r} for {year}")


def _normalise_cultivar(value: Any) -> str:
    return str(value or "").strip().casefold().replace("ё", "е")


def load_cultivar_observations(config: dict[str, Any]) -> list[CalibrationObservation]:
    if load_workbook is None:
        raise CalibrationError("openpyxl is required to read the official observation workbook")
    calibration = config["calibration"]
    root = config["_root"]
    workbook_path = root / calibration["observations_workbook"]
    if not workbook_path.is_file():
        raise CalibrationError(f"Observation workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet_name = str(calibration["observations_sheet"])
    if sheet_name not in workbook.sheetnames:
        raise CalibrationError(f"Observation sheet {sheet_name!r} not found; available={workbook.sheetnames}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: position for position, name in enumerate(headers)}

    required = {
        "Год", "Поле", "Культура", "Сорт", "Технология", "Средняя урожайность, т/га",
        "Дата посева", "Дата уборки", "Фенофаза колошения, дата начала",
    }
    missing = required - index.keys()
    if missing:
        raise CalibrationError(f"Observation workbook lacks columns: {sorted(missing)}")

    requested_years = {int(value) for value in calibration["years"]}
    calibration_years = {int(value) for value in calibration["calibration_years"]}
    validation_years = {int(value) for value in calibration["validation_years"]}
    target_cultivar = _normalise_cultivar(calibration["source_cultivar_name"])
    technology_filter = str(calibration.get("technology_filter", "")).casefold()
    selected: list[CalibrationObservation] = []
    for row in rows:
        year_number = _number(row[index["Год"]])
        if year_number is None or int(year_number) not in requested_years:
            continue
        year = int(year_number)
        if _normalise_cultivar(row[index["Сорт"]]) != target_cultivar:
            continue
        technology = str(row[index["Технология"]] or "")
        if technology_filter and technology_filter not in technology.casefold():
            continue
        crop = str(row[index["Культура"]] or "").casefold()
        if "пшениц" not in crop:
            continue

        sowing = _date_from_excel(row[index["Дата посева"]], year)
        heading = _date_from_excel(row[index["Фенофаза колошения, дата начала"]], year)
        if sowing is None or heading is None:
            raise CalibrationError(f"Year {year} lacks exact sowing or heading date")
        density = _number(row[index.get("Плотность посева, млн. зерен/га", -1)]) if "Плотность посева, млн. зерен/га" in index else None
        germination = _number(row[index.get("Всхожесть, %", -1)]) if "Всхожесть, %" in index else None
        if density is None or germination is None:
            raise CalibrationError(f"Year {year} lacks seed density or germination for population estimate")
        population = density * 100.0 * germination / 100.0
        role = "calibration" if year in calibration_years else "validation" if year in validation_years else "unused"
        selected.append(
            CalibrationObservation(
                year=year,
                field=int(float(row[index["Поле"]])),
                cultivar=str(row[index["Сорт"]]),
                technology=technology,
                observed_yield_t_ha=float(row[index["Средняя урожайность, т/га"]]),
                sowing_date=sowing,
                heading_start_date=heading,
                heading_end_date=_date_from_excel(row[index["Фенофаза колошения, дата завершения"]], year)
                if "Фенофаза колошения, дата завершения" in index else None,
                harvest_date=_date_from_excel(row[index["Дата уборки"]], year),
                thousand_seed_weight_g=_number(row[index["Масса 1000 зёрен, г"]])
                if "Масса 1000 зёрен, г" in index else None,
                sowing_rate_kg_ha=_number(row[index["Норма высева, кг/га"]])
                if "Норма высева, кг/га" in index else None,
                germination_percent=germination,
                seed_density_million_ha=density,
                population_plants_m2=population,
                dataset_role=role,
            )
        )
    workbook.close()
    selected.sort(key=lambda item: item.year)
    years = [item.year for item in selected]
    if years != sorted(requested_years):
        raise CalibrationError(f"Expected one filtered observation for years {sorted(requested_years)}, found {years}")
    return selected


# Backward-compatible name retained for existing Zlata configs and tests.
load_zlata_observations = load_cultivar_observations


def parameter_combinations(config: dict[str, Any]) -> list[dict[str, float]]:
    grid = config["calibration"].get("parameter_grid", {})
    if not grid:
        return [{}]
    names = list(grid)
    values = []
    for name in names:
        axis = grid[name].get("values", [])
        if not axis:
            raise CalibrationError(f"Calibration parameter {name} has no values")
        values.append([float(value) for value in axis])
    return [dict(zip(names, combination)) for combination in itertools.product(*values)]


def calibration_candidate_id(parameters: dict[str, float]) -> str:
    payload = json.dumps(parameters, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:12]


def _date_to_manager_value(value: date) -> str:
    return f"{value.day}-{value.strftime('%b').lower()}"


def _weather_value(calibration: dict[str, Any], year: int) -> str:
    weather_files = calibration["weather_files"]
    value = weather_files.get(year, weather_files.get(str(year)))
    if not value:
        raise CalibrationError(f"No calibration weather file configured for {year}")
    return str(value)


def _calibration_simulation_start(calibration: dict[str, Any], observation: CalibrationObservation) -> date:
    days_before = calibration.get("simulation_start_days_before_sowing")
    if days_before is None:
        return date(observation.year, 1, 1)
    days = int(days_before)
    if days < 0:
        raise CalibrationError("simulation_start_days_before_sowing must be zero or positive")
    return observation.sowing_date - timedelta(days=days)


def _patch_year_specific_initial_water(
    data: dict[str, Any], calibration: dict[str, Any], year: int
) -> list[float] | None:
    """Apply a reviewed AirDry-to-DUL fraction only to a temporary copy."""
    configured = calibration.get("initial_water_airdry_to_dul_fraction_by_year")
    if not configured:
        return None
    raw_fraction = configured.get(year, configured.get(str(year)))
    if raw_fraction is None:
        raise CalibrationError(f"No initial-water fraction configured for calibration year {year}")
    fraction = float(raw_fraction)
    if not 0.0 <= fraction <= 1.0:
        raise CalibrationError(f"Initial-water fraction for {year} must be between 0 and 1")
    layers = int(calibration.get("initial_water_layers_to_adjust", 2))
    water = _named(data, "Water")
    physical = _named(data, "Physical")
    values = list(water.get("InitialValues") or [])
    air_dry = list(physical.get("AirDry") or [])
    dul = list(physical.get("DUL") or [])
    if not values or len(values) != len(air_dry) or len(values) != len(dul):
        raise CalibrationError("Water.InitialValues, Physical.AirDry and Physical.DUL layer counts differ")
    if layers < 1 or layers > len(values):
        raise CalibrationError(f"initial_water_layers_to_adjust must be in 1..{len(values)}")
    for index in range(layers):
        values[index] = float(air_dry[index]) + fraction * (float(dul[index]) - float(air_dry[index]))
    water["InitialValues"] = values
    return values


def _patch_cultivar_parameters(data: dict[str, Any], cultivar: str, config: dict[str, Any], values: dict[str, float]) -> None:
    cultivar_node = _named(data, cultivar)
    commands = list(cultivar_node.get("Command", []))
    mapping = config["calibration"].get("parameter_grid", {})
    for parameter_name, value in values.items():
        parameter_config = mapping[parameter_name]
        command_path = str(parameter_config["command"])
        matches = [index for index, command in enumerate(commands) if str(command).split("=", 1)[0].strip() == command_path]
        if len(matches) == 0 and bool(parameter_config.get("allow_new_command", False)):
            commands.append(f"{command_path} = {value:g}")
            continue
        if len(matches) != 1:
            raise CalibrationError(
                f"Expected one existing cultivar command for {parameter_name} ({command_path}), found {len(matches)}. "
                "Set allow_new_command: true only for a verified APSIMX cultivar override path."
            )
        commands[matches[0]] = f"{command_path} = {value:g}"
    cultivar_node["Command"] = commands


def _ensure_calibration_cultivar(data: dict[str, Any], calibration: dict[str, Any]) -> None:
    """Clone a configured source cultivar in temporary/calibration candidates only."""
    target = str(calibration["cultivar"])
    matches = [node for node in _walk(data) if node.get("Name") == target]
    if len(matches) == 1:
        return
    if matches:
        raise CalibrationError(f"Expected at most one cultivar named {target!r}; found {len(matches)}")
    source_name = str(calibration.get("source_template_cultivar", ""))
    if not source_name:
        raise CalibrationError(
            f"Cultivar {target!r} is absent from template and source_template_cultivar is not configured"
        )
    source = _named(data, source_name)
    if "Cultivar" not in str(source.get("$type", "")):
        raise CalibrationError(f"Source node {source_name!r} is not an APSIMX cultivar")
    parents = [
        node for node in _walk(data)
        if isinstance(node.get("Children"), list) and source in node["Children"]
    ]
    if len(parents) != 1:
        raise CalibrationError(f"Expected one parent for source cultivar {source_name!r}; found {len(parents)}")
    cloned = copy.deepcopy(source)
    cloned["Name"] = target
    cloned["ResourceName"] = None
    parents[0]["Children"].append(cloned)


def create_calibration_copy(
    optimizer_config: dict[str, Any],
    calibration_config: dict[str, Any],
    observation: CalibrationObservation,
    parameters: dict[str, float],
    destination: Path,
) -> Path:
    calibration = calibration_config["calibration"]
    root = calibration_config["_root"]
    template = root / calibration["base_template"]
    weather = root / _weather_value(calibration, observation.year)
    if not template.is_file() or not weather.is_file():
        raise CalibrationError(f"Missing calibration template or weather: {template}, {weather}")
    fertiliser = calibration["fertilizer"]
    scheme = {
        "id": "CAL_N30",
        "passes": 1,
        "events": [{"type_apsimx": fertiliser["type_apsimx"], "amount": fertiliser["amount_kg_ha"]}],
    }
    fixed = FixedParameters(
        _date_to_manager_value(observation.sowing_date),
        observation.population_plants_m2,
        float(calibration["sowing_depth_mm"]),
        float(calibration["row_spacing_mm"]),
        [],
    )
    per_run_config = dict(optimizer_config)
    simulation_start = _calibration_simulation_start(calibration, observation)
    per_run_config["simulation"] = {
        "start_date": simulation_start.isoformat(),
        "end_date": f"{observation.year}-12-31",
    }
    create_working_copy(template, destination, weather, calibration["cultivar"], scheme, per_run_config, fixed)
    data = json.loads(destination.read_text(encoding="utf-8"))
    _ensure_calibration_cultivar(data, calibration)
    # Exact observed sowing dates are calibration inputs. Disable the weather
    # thresholds only in the temporary copy so that sowing cannot silently be
    # skipped on the one-day window.
    sowing = _parameters(_named(data, optimizer_config["apsim"]["manager_names"]["sowing"]))
    for key in ("MinESW", "MinRain"):
        if key not in sowing:
            raise CalibrationError(f"Sowing manager lacks {key}; exact-date calibration cannot be guaranteed")
        sowing[key]["Value"] = "-1"
    _patch_cultivar_parameters(data, calibration["cultivar"], calibration_config, parameters)
    _patch_year_specific_initial_water(data, calibration, observation.year)
    reference_height = calibration.get("microclimate_reference_height_m")
    if reference_height is not None:
        microclimate = _named(data, "MicroClimate")
        microclimate["ReferenceHeight"] = float(reference_height)
    destination.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return destination


def write_calibrated_candidate_template(
    calibration_config: dict[str, Any], parameters: dict[str, float], destination: Path
) -> Path:
    """Write a reviewable candidate without changing the source template."""
    calibration = calibration_config["calibration"]
    source = calibration_config["_root"] / calibration["base_template"]
    data = json.loads(source.read_text(encoding="utf-8"))
    _ensure_calibration_cultivar(data, calibration)
    _patch_cultivar_parameters(data, calibration["cultivar"], calibration_config, parameters)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return destination


def build_calibration_plan(
    optimizer_config: dict[str, Any], calibration_config: dict[str, Any], stage_files: bool = True
) -> list[dict[str, Any]]:
    observations = load_cultivar_observations(calibration_config)
    combinations = parameter_combinations(calibration_config)
    root = calibration_config["_root"]
    work_dir = root / calibration_config["calibration"]["work_dir"]
    plan = []
    for combination in combinations:
        candidate_id = calibration_candidate_id(combination)
        for observation in observations:
            cultivar_slug = str(calibration_config["calibration"]["cultivar"]).strip().casefold().replace(" ", "_")
            destination = work_dir / candidate_id / f"{cultivar_slug}_{observation.year}.apsimx"
            if stage_files:
                create_calibration_copy(optimizer_config, calibration_config, observation, combination, destination)
            row = {
                "candidate_id": candidate_id,
                "year": observation.year,
                "field": observation.field,
                "dataset_role": observation.dataset_role,
                "cultivar": calibration_config["calibration"]["cultivar"],
                "technology": observation.technology,
                "weather_file": _weather_value(calibration_config["calibration"], observation.year),
                "working_apsimx": destination.relative_to(root).as_posix(),
                "sowing_date": observation.sowing_date.isoformat(),
                "simulation_start_date": _calibration_simulation_start(
                    calibration_config["calibration"], observation
                ).isoformat(),
                "observed_heading_date": observation.heading_start_date.isoformat(),
                "observed_harvest_date": observation.harvest_date.isoformat() if observation.harvest_date else "",
                "observed_yield_t_ha": observation.observed_yield_t_ha,
                "population_plants_m2": round(observation.population_plants_m2, 3),
                "population_status": "estimated_from_seed_density_and_laboratory_germination",
                "fertilizer_type": calibration_config["calibration"]["fertilizer"]["type_apsimx"],
                "fertilizer_amount_kg_ha": calibration_config["calibration"]["fertilizer"]["amount_kg_ha"],
                "fertilizer_status": calibration_config["calibration"]["fertilizer"]["status"],
            }
            row.update(combination)
            plan.append(row)
    return plan


def _table_columns(connection: sqlite3.Connection, table: str) -> list[str]:
    return [str(row[1]) for row in connection.execute(f'PRAGMA table_info("{table}")')]


def _parse_sqlite_date(value: Any) -> date | None:
    if value is None:
        return None
    text = str(value).strip()
    for candidate in (text, text.split(" ", 1)[0]):
        try:
            return datetime.fromisoformat(candidate).date()
        except ValueError:
            pass
    return None


def read_phenology_result(database: Path, calibration_config: dict[str, Any]) -> dict[str, Any]:
    reporter = calibration_config["calibration"]["reporter"]
    stage_column = reporter["stage_column"]
    connection = sqlite3.connect(database)
    try:
        connection.row_factory = sqlite3.Row
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "Report" not in tables:
            raise CalibrationError(f"Report table is absent in {database}; tables={sorted(tables)}")
        columns = _table_columns(connection, "Report")
        if stage_column not in columns:
            raise CalibrationError(f"Phenology column {stage_column!r} is absent; columns={columns}")
        date_column = next((name for name in reporter["date_candidates"] if name in columns), None)
        if date_column is None:
            raise CalibrationError(f"No exact report date column found; candidates={reporter['date_candidates']}, columns={columns}")
        yield_column = next((name for name in ("Yield", "GrainWt", "Wheat.Yield") if name in columns), None)
        zadok_column = str(reporter.get("zadok_column", ""))
        if zadok_column not in columns:
            zadok_column = ""
        selected = [date_column, stage_column]
        if zadok_column:
            selected.append(zadok_column)
        if yield_column:
            selected.append(yield_column)
        rows = connection.execute(
            "SELECT " + ", ".join(f'"{name}"' for name in selected) + ' FROM "Report" ORDER BY rowid'
        ).fetchall()
    finally:
        connection.close()

    def first_stage(names: list[str]) -> tuple[date | None, str]:
        wanted = {name.casefold() for name in names}
        for row in rows:
            stage_name = str(row[stage_column] or "").strip()
            if stage_name.casefold() in wanted:
                parsed = _parse_sqlite_date(row[date_column])
                if parsed:
                    return parsed, stage_name
        return None, ""

    heading_date, heading_stage = first_stage(list(reporter["heading_stage_names"]))
    heading_method = "current_stage_name" if heading_date else ""
    if heading_date is None and zadok_column:
        threshold = float(reporter.get("heading_zadok_min", 50))
        for row in rows:
            zadok = _number(row[zadok_column])
            if zadok is not None and zadok >= threshold:
                heading_date = _parse_sqlite_date(row[date_column])
                if heading_date:
                    heading_method = f"zadok_gte_{threshold:g}"
                    heading_stage = ""
                    break

    yield_raw = None
    if yield_column:
        values = [_number(row[yield_column]) for row in rows]
        usable = [value for value in values if value is not None and value > 0]
        yield_raw = usable[-1] if usable else 0.0
    maturity_date, maturity_stage = first_stage(list(reporter["maturity_stage_names"]))
    return {
        "simulated_heading_date": heading_date,
        "simulated_heading_stage": heading_stage,
        "heading_detection_method": heading_method,
        "simulated_maturity_date": maturity_date,
        "simulated_maturity_stage": maturity_stage,
        "simulated_yield_raw": yield_raw,
        "yield_column": yield_column,
        "available_columns": columns,
    }


def maturity_upper_bound_metrics(simulated_maturity: date | None, observed_harvest: date | None) -> dict[str, Any]:
    """Compare simulated maturity with harvest used only as an upper bound."""
    if observed_harvest is None:
        return {
            "maturity_status": "harvest_date_missing",
            "maturity_error_days": None,
            "maturity_late_days": None,
            "maturity_slack_days": None,
        }
    if simulated_maturity is None:
        return {
            "maturity_status": "maturity_not_reached",
            "maturity_error_days": None,
            "maturity_late_days": None,
            "maturity_slack_days": None,
        }
    error_days = (simulated_maturity - observed_harvest).days
    return {
        "maturity_status": "late" if error_days > 0 else "on_or_before_harvest",
        "maturity_error_days": error_days,
        "maturity_late_days": max(0, error_days),
        "maturity_slack_days": max(0, -error_days),
    }


def run_calibration_plan(
    plan: list[dict[str, Any]], optimizer_config: dict[str, Any], calibration_config: dict[str, Any], logger
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    root = calibration_config["_root"]
    run_rows = []
    failed_candidates: dict[str, str] = {}
    for row in plan:
        working = root / row["working_apsimx"]
        result = dict(row)
        candidate_id = row["candidate_id"]
        if candidate_id in failed_candidates:
            result.update(
                {
                    "simulated_heading_date": "",
                    "simulated_heading_stage": "",
                    "heading_detection_method": "",
                    "heading_error_days": None,
                    "simulated_maturity_date": "",
                    "simulated_maturity_stage": "",
                    "maturity_status": "run_skipped",
                    "maturity_error_days": None,
                    "maturity_late_days": None,
                    "maturity_slack_days": None,
                    "simulated_yield_raw": None,
                    "run_status": "skipped_after_candidate_failure",
                    "error_message": (
                        "Candidate was not run after an earlier season failed: "
                        f"{failed_candidates[candidate_id]}"
                    ),
                }
            )
            logger.warning(
                "Cultivar calibration run skipped after candidate failure: candidate=%s year=%s",
                candidate_id,
                row["year"],
            )
            run_rows.append(result)
            continue
        try:
            run_apsim_files([working], optimizer_config, logger)
            database = find_db_for_apsimx(working)
            if database is None:
                raise CalibrationError(f"APSIMX database not found for {working}")
            phenology = read_phenology_result(database, calibration_config)
            simulated = phenology["simulated_heading_date"]
            observed = date.fromisoformat(row["observed_heading_date"])
            maturity = phenology["simulated_maturity_date"]
            harvest = date.fromisoformat(row["observed_harvest_date"]) if row.get("observed_harvest_date") else None
            maturity_metrics = maturity_upper_bound_metrics(maturity, harvest)
            result.update(
                {
                    "simulated_heading_date": simulated.isoformat() if simulated else "",
                    "simulated_heading_stage": phenology["simulated_heading_stage"],
                    "heading_detection_method": phenology["heading_detection_method"],
                    "heading_error_days": (simulated - observed).days if simulated else None,
                    "simulated_maturity_date": maturity.isoformat() if maturity else "",
                    "simulated_maturity_stage": phenology["simulated_maturity_stage"],
                    "simulated_yield_raw": phenology["simulated_yield_raw"],
                    "run_status": "completed" if simulated else "completed_without_heading_stage",
                    "error_message": "" if simulated else "Configured heading stage was not found in Report",
                }
            )
            result.update(maturity_metrics)
        except Exception as exc:
            failed_candidates[candidate_id] = str(exc)
            result.update(
                {
                    "simulated_heading_date": "",
                    "simulated_heading_stage": "",
                    "heading_detection_method": "",
                    "heading_error_days": None,
                    "simulated_maturity_date": "",
                    "simulated_maturity_stage": "",
                    "maturity_status": "run_failed",
                    "maturity_error_days": None,
                    "maturity_late_days": None,
                    "maturity_slack_days": None,
                    "simulated_yield_raw": None,
                    "run_status": "failed",
                    "error_message": str(exc),
                }
            )
            logger.exception("Cultivar calibration run failed: candidate=%s year=%s", row["candidate_id"], row["year"])
        run_rows.append(result)

    summaries = []
    for candidate_id in dict.fromkeys(row["candidate_id"] for row in run_rows):
        group = [row for row in run_rows if row["candidate_id"] == candidate_id]
        parameter_names = list(calibration_config["calibration"].get("parameter_grid", {}))
        summary = {key: group[0][key] for key in parameter_names}
        summary.update({"candidate_id": candidate_id})
        for role in ("calibration", "validation"):
            role_rows = [row for row in group if row["dataset_role"] == role]
            errors = [float(row["heading_error_days"]) for row in role_rows if row["heading_error_days"] is not None]
            summary[f"{role}_n"] = len(errors)
            summary[f"{role}_heading_mae_days"] = statistics_fmean_abs(errors) if errors else None
            summary[f"{role}_heading_rmse_days"] = math.sqrt(sum(value * value for value in errors) / len(errors)) if errors else None
            late_days = [float(row["maturity_late_days"]) for row in role_rows if row.get("maturity_late_days") is not None]
            slack_days = [float(row["maturity_slack_days"]) for row in role_rows if row.get("maturity_slack_days") is not None]
            summary[f"{role}_maturity_n"] = len(late_days)
            summary[f"{role}_maturity_missing_count"] = sum(
                row.get("maturity_status") == "maturity_not_reached" for row in role_rows
            )
            summary[f"{role}_maturity_on_time_count"] = sum(
                row.get("maturity_status") == "on_or_before_harvest" for row in role_rows
            )
            summary[f"{role}_maturity_late_rmse_days"] = (
                math.sqrt(sum(value * value for value in late_days) / len(late_days)) if late_days else None
            )
            summary[f"{role}_maturity_max_late_days"] = max(late_days) if late_days else None
            summary[f"{role}_maturity_mean_slack_days"] = sum(slack_days) / len(slack_days) if slack_days else None
        summary["actual_failed_runs"] = sum(row["run_status"] == "failed" for row in group)
        summary["skipped_runs"] = sum(
            row["run_status"] == "skipped_after_candidate_failure" for row in group
        )
        summary["failed_runs"] = summary["actual_failed_runs"] + summary["skipped_runs"]
        summary["run_status"] = "completed" if all(row["run_status"] == "completed" for row in group) else "partial_or_failed"
        summaries.append(summary)
    selection_mode = str(
        calibration_config["calibration"].get("selection", {}).get("mode", "heading_rmse")
    )
    if selection_mode == "maturity_upper_bound_then_heading":
        summaries.sort(
            key=lambda row: (
                row["failed_runs"],
                row["calibration_maturity_missing_count"],
                _sort_number(row["calibration_maturity_late_rmse_days"]),
                _sort_number(row["calibration_heading_rmse_days"]),
                _sort_number(row["calibration_maturity_mean_slack_days"]),
                row["candidate_id"],
            )
        )
    elif selection_mode == "heading_rmse":
        summaries.sort(
            key=lambda row: (
                row["failed_runs"],
                row["calibration_heading_rmse_days"] is None,
                _sort_number(row["calibration_heading_rmse_days"]),
                row["candidate_id"],
            )
        )
    else:
        raise CalibrationError(f"Unknown calibration selection mode: {selection_mode}")
    return run_rows, summaries


DIAGNOSTIC_NUMERIC_COLUMNS = (
    "Wheat.LAI",
    "Wheat.AboveGround.Wt",
    "Wheat.AboveGround.N",
    "Wheat.Grain.Total.Wt",
    "Wheat.Grain.Total.N",
    "Wheat.Grain.Number",
    "Yield",
    "PAW",
    "WaterStored",
    "MineralN",
    "MineralNWithUrea",
    "NO3kgHa",
    "NH4kgHa",
    "Drainage",
    "Runoff",
    "PondDepth",
    "LeafFw",
    "LeafFn",
    "LeafLAIDead",
    "LeafWaterDemand",
    "LeafWaterAllocation",
    "LeafPotentialEP",
    "LeafTranspiration",
    "LeafPhotosynthesis",
    "LeafRadiationIntercepted",
    "RootDepth",
    "RootFASW",
    "RootWaterTensionFactor",
    "RootWaterUptake",
    "RootNUptake",
)


def add_calibration_diagnostic_report_expressions(
    apsimx_path: Path, calibration_config: dict[str, Any]
) -> list[str]:
    """Add explicitly configured, source-verified expressions to a temporary copy."""
    expressions = list(calibration_config["calibration"].get("diagnostic_report_expressions", []))
    if not expressions:
        return []
    data = json.loads(apsimx_path.read_text(encoding="utf-8"))
    reports = [
        node for node in _walk(data)
        if node.get("Name") == "Report" and "Models.Report" in str(node.get("$type", ""))
    ]
    if len(reports) != 1:
        raise CalibrationError(f"Expected exactly one Report node for diagnostics, found {len(reports)}")
    variable_names = reports[0].setdefault("VariableNames", [])
    added = []
    for expression in expressions:
        if expression not in variable_names:
            variable_names.append(expression)
            added.append(expression)
    apsimx_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return added


def read_calibration_daily_report(
    database: Path, calibration_config: dict[str, Any], plan_row: dict[str, Any]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Export exact existing Report columns without inventing APSIM expressions."""
    reporter = calibration_config["calibration"]["reporter"]
    connection = sqlite3.connect(database)
    try:
        connection.row_factory = sqlite3.Row
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "Report" not in tables:
            raise CalibrationError(f"Report table is absent in {database}; tables={sorted(tables)}")
        columns = _table_columns(connection, "Report")
        rows = connection.execute('SELECT rowid AS "report_rowid", * FROM "Report" ORDER BY rowid').fetchall()
    finally:
        connection.close()

    exported: list[dict[str, Any]] = []
    for source in rows:
        item = {
            "candidate_id": plan_row["candidate_id"],
            "year": plan_row["year"],
            "dataset_role": plan_row["dataset_role"],
        }
        item.update({key: source[key] for key in source.keys()})
        exported.append(item)

    date_column = next((name for name in reporter["date_candidates"] if name in columns), None)
    stage_column = str(reporter["stage_column"])
    summary: dict[str, Any] = {
        "candidate_id": plan_row["candidate_id"],
        "year": plan_row["year"],
        "dataset_role": plan_row["dataset_role"],
        "report_row_count": len(exported),
        "report_start_date": exported[0].get(date_column, "") if exported and date_column else "",
        "report_end_date": exported[-1].get(date_column, "") if exported and date_column else "",
        "last_crop_stage": exported[-1].get(stage_column, "") if exported and stage_column in columns else "",
        "available_report_columns": " | ".join(columns),
        "direct_water_stress_index_status": "not_reported",
        "direct_nitrogen_stress_index_status": "not_reported",
        "water_diagnostics_status": "available" if {"PAW", "WaterStored"} & set(columns) else "missing",
        "nitrogen_diagnostics_status": (
            "available" if "MineralN" in columns or {"NO3kgHa", "NH4kgHa"} <= set(columns) else "missing"
        ),
    }
    for column in DIAGNOSTIC_NUMERIC_COLUMNS:
        if column not in columns:
            summary[f"{column}_status"] = "missing"
            continue
        values = [_number(row.get(column)) for row in exported]
        usable = [value for value in values if value is not None]
        summary[f"{column}_status"] = "available" if usable else "all_null"
        summary[f"{column}_min"] = min(usable) if usable else None
        summary[f"{column}_max"] = max(usable) if usable else None
        summary[f"{column}_last"] = usable[-1] if usable else None
        if column in {"Drainage", "Runoff"}:
            summary[f"{column}_sum"] = sum(usable) if usable else None
    for prefix in ("PAW", "SW"):
        layer_columns = [column for column in columns if column.startswith(f"{prefix}(")]
        if not layer_columns:
            continue
        totals = []
        top_values = []
        for row in exported:
            values = [_number(row.get(column)) for column in layer_columns]
            if all(value is not None for value in values):
                totals.append(sum(values))
                top_values.append(values[0])
        summary[f"{prefix}_status"] = "available" if totals else "all_null"
        summary[f"{prefix}_layer_count"] = len(layer_columns)
        summary[f"{prefix}_total_min"] = min(totals) if totals else None
        summary[f"{prefix}_total_max"] = max(totals) if totals else None
        summary[f"{prefix}_total_last"] = totals[-1] if totals else None
        summary[f"{prefix}_top_layer_min"] = min(top_values) if top_values else None
        summary[f"{prefix}_top_layer_last"] = top_values[-1] if top_values else None
    yield_max = summary.get("Yield_max")
    if yield_max is not None:
        summary["yield_max_t_ha"], summary["yield_conversion_status"] = convert_yield_to_t_ha(
            yield_max, reporter
        )
    else:
        summary["yield_max_t_ha"] = None
        summary["yield_conversion_status"] = "missing"
    return exported, summary


def collect_calibration_diagnostics(
    plan: list[dict[str, Any]], calibration_config: dict[str, Any]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    root = calibration_config["_root"]
    daily_rows: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    for row in plan:
        working = root / row["working_apsimx"]
        database = find_db_for_apsimx(working)
        if database is None:
            raise CalibrationError(f"APSIMX database not found for diagnostic run {working}")
        daily, summary = read_calibration_daily_report(database, calibration_config, row)
        daily_rows.extend(daily)
        summaries.append(summary)
    return daily_rows, summaries


def _sort_number(value: Any) -> float:
    return float(value) if value is not None else float("inf")


def statistics_fmean_abs(values: list[float]) -> float:
    return sum(abs(value) for value in values) / len(values)


def write_rows(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(dict.fromkeys(key for row in rows for key in row)) if rows else ["candidate_id"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
